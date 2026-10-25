from __future__ import annotations

import copy
import datetime as dt
import hashlib
import io
import json
import tempfile
import unittest
from collections import Counter
from pathlib import Path
from unittest import mock

import duckdb
import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import yaml

import nhse_procurement
import nhse_spending
import nhse_spending_inventory


HEADERS = [
    "Department family",
    "Entity",
    "Date",
    "Expense Type for Report",
    "Expense Area for Report",
    "Supplier",
    "Transaction number",
    "AP Amount (£)",
    "VAT registration number",
]

HEADERS_V2 = [
    "Department family",
    "Entity",
    "Date",
    "Expense Type",
    "Expense Area",
    "Supplier",
    "Transaction number",
    "AP Amount (£)",
    "VAT registration number",
    "Purchase invoice number",
]


class SpendingDataTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def _workbook(self, headers: list[str] = HEADERS) -> Path:
        path = self.root / "fixture.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(headers)
        sheet.append(["Department of Health", "NHS England", dt.date(2025, 9, 1), "Services", "Digital", "Supplier A", 1001, 10.254, "123456789"])
        sheet.append(["Department of Health", "NHS England", dt.date(2025, 9, 2), "Services", "Digital", "Removed", "1002", -2.25, None])
        sheet.append(["Department of Health", "North of England CSU", dt.date(2025, 9, 3), "Goods", "Operations", "Supplier B", "1003", 5, None])
        workbook.save(path)
        return path

    def _source(self, path: Path) -> dict:
        return {
            "id": "fixture",
            "title": "Fixture payments",
            "normalization_schema_version": 1,
            "period": "2025-09",
            "publication_date": dt.date(2025, 10, 1),
            "publication_date_precision": "month",
            "publisher": "NHS England",
            "landing_page_url": "https://example.test/landing",
            "landing_page_archive_url": "https://example.test/landing/archive",
            "collection_url": "https://example.test/collection",
            "collection_archive_url": "https://example.test/collection/archive",
            "source_url": "https://example.test/source.xlsx",
            "archive_url": "https://example.test/source/archive.xlsx",
            "filename": path.name,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "bytes": path.stat().st_size,
            "rights_status": "test-only",
            "rights_basis": "test fixture",
            "rights_decision_date": dt.date(2025, 10, 1),
            "rights_note": "Fixture data",
            "rights_evidence_url": "https://example.test/rights",
            "rights_evidence_archive_url": "https://example.test/rights/archive",
            "referenced_licence_url": "https://example.test/licence",
            "referenced_licence_archive_url": "https://example.test/licence/archive",
            "workbook": {"sheet": "Payments", "amount_scale": 2, "columns": HEADERS},
            "expected": {
                "rows": 3,
                "total_gbp": "13.00",
                "nhs_england_rows": 2,
                "nhs_england_total_gbp": "8.00",
                "negative_rows": 1,
                "missing_vat_rows": 2,
            },
        }

    def test_build_is_deterministic_and_queryable(self) -> None:
        path = self._workbook()
        source = self._source(path)
        first = nhse_spending.build_release(path, source, self.root / "first")
        second = nhse_spending.build_release(path, source, self.root / "second")

        first_files = {item.name: item.read_bytes() for item in first.iterdir()}
        second_files = {item.name: item.read_bytes() for item in second.iterdir()}
        self.assertEqual(first_files, second_files)

        connection = duckdb.connect(":memory:")
        try:
            result = connection.execute(
                """
                SELECT count(*), sum(amount_gbp), min(source_row), max(source_row),
                       min(source_amount_raw),
                       max(source_amount_raw) FILTER (WHERE transaction_number = '1001')
                FROM read_parquet(?)
                """,
                [str(first / "payments.parquet")],
            ).fetchone()
        finally:
            connection.close()
        self.assertEqual(result, (3, nhse_spending.Decimal("13.00"), 2, 4, "-2.25", "10.254"))

        validation = json.loads((first / "validation.json").read_text(encoding="utf-8"))
        self.assertEqual(validation["statistics"]["entity_counts"]["NHS England"], 2)
        self.assertEqual(validation["statistics"]["entity_totals_gbp"]["NHS England"], "8.00")

    def test_build_release_round_trips_no_transaction_number_column(self) -> None:
        headers_no_transaction_number = [
            "Department family", "Entity", "Date", "Expense Type", "Expense Area",
            "Supplier", "AP Amount (£)", "VAT registration number", "Purchase invoice number",
        ]
        path = self.root / "no-transaction-number-build.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(headers_no_transaction_number)
        sheet.append(["Department of Health", "NHS England", dt.date(2019, 5, 1), "Services", "Digital", "Supplier A", 10.25, "123456789", "INV-1"])
        workbook.save(path)
        source = self._source(path)
        source["normalization_schema_version"] = 2
        source["workbook"] = {
            "sheet": "Payments",
            "amount_scale": 2,
            "columns": headers_no_transaction_number,
            "no_transaction_number_column": True,
        }
        source["expected"] = {
            "rows": 1,
            "total_gbp": "10.25",
            "nhs_england_rows": 1,
            "nhs_england_total_gbp": "10.25",
            "negative_rows": 0,
            "missing_vat_rows": 0,
        }
        target = nhse_spending.build_release(path, source, self.root / "no-transaction-number-release")
        nhse_spending.validate_release(target, source, verify_against_sources=True, source_files=[path])

        table = pq.read_table(target / "payments.parquet")
        self.assertTrue(table.schema.field("transaction_number").nullable)
        self.assertIsNone(table.to_pylist()[0]["transaction_number"])

        # A workbook with no actual null transaction_number keeps the field non-nullable,
        # so already-published releases are not retroactively re-encoded.
        normal_path = self._workbook()
        normal_source = self._source(normal_path)
        normal_target = nhse_spending.build_release(normal_path, normal_source, self.root / "normal-release")
        normal_table = pq.read_table(normal_target / "payments.parquet")
        self.assertFalse(normal_table.schema.field("transaction_number").nullable)

    def test_schema_drift_fails_closed(self) -> None:
        wrong_headers = copy.copy(HEADERS)
        wrong_headers[7] = "Amount"
        path = self._workbook(wrong_headers)
        source = self._source(path)
        with self.assertRaisesRegex(nhse_spending.SpendingDataError, "unexpected workbook headers"):
            nhse_spending.build_release(path, source, self.root / "release")

    def test_trailing_columns_fail_closed_by_default(self) -> None:
        path = self.root / "trailing.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(HEADERS)
        sheet.append(["Department of Health", "NHS England", dt.date(2025, 9, 1), "Services", "Digital", "Supplier A", 1001, 10.25, "123456789"])
        sheet.cell(row=2, column=10, value="unexpected")
        workbook.save(path)
        source = self._source(path)
        with self.assertRaisesRegex(nhse_spending.SpendingDataError, "ignore_trailing_columns"):
            nhse_spending.parse_workbook(path, source)

    def test_non_integral_purchase_invoice_number_fails_closed_by_default(self) -> None:
        path = self.root / "non-integral.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(HEADERS_V2)
        sheet.append(
            [
                "Department of Health", "NHS England", dt.date(2025, 9, 1), "Services", "Digital",
                "Supplier A", 1001, 10.25, "123456789", 99479.196,
            ]
        )
        workbook.save(path)
        source = self._source(path)
        source["normalization_schema_version"] = 2
        source["workbook"] = {"sheet": "Payments", "amount_scale": 2, "columns": HEADERS_V2}
        with self.assertRaisesRegex(nhse_spending.SpendingDataError, "non-integral numeric identifier"):
            nhse_spending.parse_workbook(path, source)

    def test_non_integral_purchase_invoice_number_permitted_when_registered(self) -> None:
        path = self.root / "non-integral-ok.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(HEADERS_V2)
        sheet.append(
            [
                "Department of Health", "NHS England", dt.date(2025, 9, 1), "Services", "Digital",
                "Supplier A", 1001, 10.25, "123456789", 99479.196,
            ]
        )
        workbook.save(path)
        source = self._source(path)
        source["normalization_schema_version"] = 2
        source["workbook"] = {
            "sheet": "Payments",
            "amount_scale": 2,
            "columns": HEADERS_V2,
            "permit_non_integral_purchase_invoice_numbers": True,
        }
        source["expected"] = {}
        records, stats = nhse_spending.parse_workbook(path, source)
        self.assertEqual(records[0]["purchase_invoice_number"], "99479.196")

    def test_no_transaction_number_column_parses_shifted_fields(self) -> None:
        headers_no_transaction_number = [
            "Department family", "Entity", "Date", "Expense Type", "Expense Area",
            "Supplier", "AP Amount (£)", "VAT registration number", "Purchase invoice number",
        ]
        path = self.root / "no-transaction-number.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(headers_no_transaction_number)
        sheet.append(
            ["Department of Health", "NHS England", dt.date(2019, 5, 1), "Services", "Digital", "Supplier A", 10.25, "123456789", "INV-1"]
        )
        workbook.save(path)
        source = self._source(path)
        source["normalization_schema_version"] = 2
        source["workbook"] = {
            "sheet": "Payments",
            "amount_scale": 2,
            "columns": headers_no_transaction_number,
            "no_transaction_number_column": True,
        }
        source["expected"] = {}
        records, stats = nhse_spending.parse_workbook(path, source)
        self.assertIsNone(records[0]["transaction_number"])
        self.assertEqual(records[0]["purchase_invoice_number"], "INV-1")
        self.assertEqual(records[0]["vat_registration_number"], "123456789")
        self.assertEqual(records[0]["amount_gbp"], nhse_spending.Decimal("10.25"))

    def test_no_transaction_number_column_requires_schema_v2(self) -> None:
        headers = ["Department family", "Entity", "Date", "Expense Type", "Expense Area", "Supplier", "AP Amount (£)", "VAT registration number"]
        path = self.root / "no-transaction-number-v1.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(headers)
        sheet.append(["Department of Health", "NHS England", dt.date(2019, 5, 1), "Services", "Digital", "Supplier A", 10.25, "123456789"])
        workbook.save(path)
        source = self._source(path)
        source["normalization_schema_version"] = 1
        source["workbook"] = {"sheet": "Payments", "amount_scale": 2, "columns": headers, "no_transaction_number_column": True}
        with self.assertRaisesRegex(nhse_spending.SpendingDataError, "require normalization schema v2"):
            nhse_spending.parse_workbook(path, source)

    def test_no_vat_registration_number_column_requires_schema_v2(self) -> None:
        headers = ["Department family", "Entity", "Date", "Expense Type", "Expense Area", "Supplier", "Transaction number", "AP Amount (£)"]
        path = self.root / "no-vat-v1.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(headers)
        sheet.append(["Department of Health", "NHS England", dt.date(2013, 9, 1), "Services", "Digital", "Supplier A", 1001, 10.25])
        workbook.save(path)
        source = self._source(path)
        source["normalization_schema_version"] = 1
        source["workbook"] = {"sheet": "Payments", "amount_scale": 2, "columns": headers, "no_vat_registration_number_column": True}
        with self.assertRaisesRegex(nhse_spending.SpendingDataError, "require normalization schema v2"):
            nhse_spending.parse_workbook(path, source)

    def test_excluded_rows_are_skipped_with_a_registered_reason(self) -> None:
        path = self.root / "excluded-rows.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(HEADERS)
        sheet.append(["Department of Health", "NHS England", dt.date(2025, 9, 1), "Services", "Digital", "Supplier A", 1001, 10.25, "123456789"])
        sheet.append(["Department of Health", "NHS England", dt.date(2025, 9, 2), "Services", "Digital", "Supplier B", 1002, "corrupted-url-value", "123456789"])
        workbook.save(path)
        source = self._source(path)
        source["workbook"]["excluded_rows"] = [3]
        source["workbook"]["excluded_rows_reason"] = "AP Amount cell contains a pasted document URL, not a numeric value"
        source["expected"] = {}
        records, stats = nhse_spending.parse_workbook(path, source)
        self.assertEqual(stats["rows"], 1)
        self.assertEqual(stats["excluded_malformed_rows"], 1)

    def test_excluded_rows_require_a_reason(self) -> None:
        path = self._workbook()
        source = self._source(path)
        source["workbook"]["excluded_rows"] = [2]
        with self.assertRaisesRegex(nhse_spending.SpendingDataError, "excluded_rows_reason"):
            nhse_spending.parse_workbook(path, source)

    def test_excluded_rows_count_must_match_rows_encountered(self) -> None:
        path = self._workbook()
        source = self._source(path)
        source["workbook"]["excluded_rows"] = [2, 9999]
        source["workbook"]["excluded_rows_reason"] = "test"
        with self.assertRaisesRegex(nhse_spending.SpendingDataError, "did not match"):
            nhse_spending.parse_workbook(path, source)

    def test_schema_v2_without_purchase_invoice_column_fails_closed(self) -> None:
        path = self._workbook()
        source = self._source(path)
        source["normalization_schema_version"] = 2
        with self.assertRaisesRegex(nhse_spending.SpendingDataError, "requires a purchase invoice number column"):
            nhse_spending.parse_workbook(path, source)

    def test_unrecognized_expected_stat_fails_closed(self) -> None:
        path = self._workbook()
        source = self._source(path)
        source["expected"]["excluded_malformed_rows"] = 1
        with self.assertRaisesRegex(nhse_spending.SpendingDataError, "not a recognized statistic"):
            nhse_spending.parse_workbook(path, source)

    def test_numeric_text_field_fails_closed_by_default(self) -> None:
        path = self.root / "numeric-text.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(HEADERS)
        sheet.append(["Department of Health", "NHS England", dt.date(2025, 9, 1), "Services", 0, "Supplier A", 1001, 10.25, "123456789"])
        workbook.save(path)
        source = self._source(path)
        with self.assertRaisesRegex(nhse_spending.SpendingDataError, "must be non-empty text"):
            nhse_spending.parse_workbook(path, source)

    def test_numeric_text_field_permitted_when_registered(self) -> None:
        path = self.root / "numeric-text-ok.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(HEADERS)
        sheet.append(["Department of Health", "NHS England", dt.date(2025, 9, 1), "Services", 0, "Supplier A", 1001, 10.25, "123456789"])
        workbook.save(path)
        source = self._source(path)
        source["workbook"]["permit_numeric_text_fields"] = True
        source["expected"] = {}
        records, stats = nhse_spending.parse_workbook(path, source)
        self.assertEqual(records[0]["expense_area"], "0")

    def test_no_vat_registration_number_column_parses_shifted_fields(self) -> None:
        headers_no_vat = [
            "Department family", "Entity", "Date", "Expense Type", "Expense Area",
            "Supplier", "Transaction number", "AP Amount (£)", "Purchase invoice number", "Comments from NHS England",
        ]
        path = self.root / "no-vat.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(headers_no_vat)
        sheet.append(
            ["Department of Health", "NHS England", dt.date(2013, 9, 1), "Services", "Digital", "Supplier A", 1001, 10.25, "INV-1", "a comment"]
        )
        workbook.save(path)
        source = self._source(path)
        source["normalization_schema_version"] = 2
        source["workbook"] = {
            "sheet": "Payments",
            "amount_scale": 2,
            "columns": headers_no_vat,
            "no_vat_registration_number_column": True,
        }
        source["expected"] = {}
        records, stats = nhse_spending.parse_workbook(path, source)
        self.assertEqual(records[0]["transaction_number"], "1001")
        self.assertIsNone(records[0]["vat_registration_number"])
        self.assertEqual(records[0]["purchase_invoice_number"], "INV-1")
        self.assertEqual(stats["missing_vat_rows"], 1)

    def test_no_transaction_number_and_no_vat_flags_are_mutually_exclusive(self) -> None:
        path = self._workbook()
        source = self._source(path)
        source["workbook"]["no_transaction_number_column"] = True
        source["workbook"]["no_vat_registration_number_column"] = True
        with self.assertRaisesRegex(nhse_spending.SpendingDataError, "mutually exclusive"):
            nhse_spending.parse_workbook(path, source)

    def test_trailing_columns_ignored_when_registered(self) -> None:
        path = self.root / "trailing-ok.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Payments"
        sheet.append(HEADERS)
        sheet.append(["Department of Health", "NHS England", dt.date(2025, 9, 1), "Services", "Digital", "Supplier A", 1001, 10.25, "123456789"])
        sheet.cell(row=2, column=10, value="internal marker")
        workbook.save(path)
        source = self._source(path)
        source["workbook"]["ignore_trailing_columns"] = True
        source["expected"] = {}
        records, stats = nhse_spending.parse_workbook(path, source)
        self.assertEqual(stats["rows"], 1)
        self.assertNotIn("internal marker", json.dumps(records[0], default=str))

    def test_multi_workbook_release_includes_revision_diff(self) -> None:
        def workbook(name: str, date: dt.date, transaction: str, vat: str | None) -> Path:
            path = self.root / name
            document = openpyxl.Workbook()
            sheet = document.active
            sheet.title = "Payments"
            sheet.append(HEADERS_V2)
            sheet.append(
                [
                    "Department of Health",
                    "NHS England",
                    date,
                    "Services",
                    "Digital",
                    "Supplier A",
                    transaction,
                    10,
                    vat,
                    f"INV-{transaction}",
                ]
            )
            document.save(path)
            return path

        june = workbook("june.xlsx", dt.date(2024, 6, 30), "1001", "123")
        original = workbook("july-original.xlsx", dt.date(2024, 6, 30), "1001", None)
        corrected = workbook("july-corrected.xlsx", dt.date(2024, 7, 31), "1002", "456")

        def entry(identifier: str, period: str, path: Path) -> dict:
            return {
                "id": identifier,
                "title": identifier,
                "period": period,
                "publication_date": dt.date(2024, 8, 1),
                "publication_date_precision": "day",
                "source_url": f"https://example.test/{path.name}",
                "archive_url": f"https://example.test/archive/{path.name}",
                "filename": path.name,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "bytes": path.stat().st_size,
                "workbook": {"sheet": "Payments", "amount_scale": 2, "columns": HEADERS_V2},
                "expected": {
                    "rows": 1,
                    "total_gbp": "10.00",
                    "nhs_england_rows": 1,
                    "nhs_england_total_gbp": "10.00",
                    "negative_rows": 0,
                    "missing_vat_rows": 1 if path == original else 0,
                },
            }

        source = {
            "id": "2024-fixture",
            "title": "2024 fixture",
            "normalization_schema_version": 2,
            "period": "2024",
            "publication_date": dt.date(2024, 8, 1),
            "publication_date_precision": "day",
            "publisher": "NHS England",
            "landing_page_url": "https://example.test/landing",
            "landing_page_archive_url": "https://example.test/landing/archive",
            "collection_url": "https://example.test/collection",
            "collection_archive_url": "https://example.test/collection/archive",
            "rights_status": "test-only",
            "rights_basis": "test fixture",
            "rights_decision_date": dt.date(2024, 8, 1),
            "rights_note": "Fixture data",
            "rights_evidence_url": "https://example.test/rights",
            "rights_evidence_archive_url": "https://example.test/rights/archive",
            "referenced_licence_url": "https://example.test/licence",
            "referenced_licence_archive_url": "https://example.test/licence/archive",
            "expected": {
                "rows": 2,
                "total_gbp": "20.00",
                "nhs_england_rows": 2,
                "nhs_england_total_gbp": "20.00",
                "negative_rows": 0,
                "missing_vat_rows": 0,
            },
            "workbooks": [
                entry("2024-06-fixture", "2024-06", june),
                entry("2024-07-fixture-corrected", "2024-07", corrected),
            ],
            "revisions": [entry("2024-07-fixture-original", "2024-07", original)],
            "revision_diff": {
                "id": "2024-07-correction",
                "original_source_id": "2024-07-fixture-original",
                "corrected_source_id": "2024-07-fixture-corrected",
                "baseline_source_id": "2024-06-fixture",
            },
        }

        with mock.patch.object(nhse_spending, "SOURCE_DIR", self.root):
            release = nhse_spending.build_release([june, corrected], source, self.root / "annual")
            second_release = nhse_spending.build_release([june, corrected], source, self.root / "annual-second")

        first_files = {item.name: item.read_bytes() for item in release.iterdir()}
        second_files = {item.name: item.read_bytes() for item in second_release.iterdir()}
        self.assertEqual(first_files, second_files)

        revision = json.loads((release / "july-2024-correction.json").read_text(encoding="utf-8"))
        self.assertEqual(revision["comparison"]["removed_rows"], 1)
        self.assertEqual(revision["comparison"]["added_rows"], 1)
        self.assertTrue(revision["diagnosis"]["ordered_rows_equal_to_baseline_excluding_vat"])

        connection = duckdb.connect(":memory:")
        try:
            result = connection.execute(
                "SELECT source_period, purchase_invoice_number FROM read_parquet(?) ORDER BY source_period",
                [str(release / "payments.parquet")],
            ).fetchall()
        finally:
            connection.close()
        self.assertEqual(result, [("2024-06", "INV-1001"), ("2024-07", "INV-1002")])

    def test_revision_fingerprint_includes_parser_visible_amount(self) -> None:
        record = {
            "department_family": "Department of Health",
            "entity": "NHS England",
            "transaction_date": dt.date(2024, 7, 1),
            "expense_type": "Services",
            "expense_area": "Digital",
            "supplier": "Supplier A",
            "transaction_number": "1001",
            "purchase_invoice_number": "INV-1001",
            "amount_gbp": nhse_spending.Decimal("10.25"),
            "source_amount_raw": "10.254",
            "currency": "GBP",
            "vat_registration_number": "123",
            "redaction_status": "none",
        }
        changed = {**record, "source_amount_raw": "10.253"}
        self.assertNotEqual(nhse_spending._semantic_row(record), nhse_spending._semantic_row(changed))

    def test_committed_releases_validate(self) -> None:
        source_ids = [f"{year}-payments-over-25000" for year in range(2014, 2025)]
        source_ids.insert(0, "2013-14-payments-over-25000")
        source_ids.append("2025-09-payments-over-25000")
        for source_id in source_ids:
            source = nhse_spending.load_source(source_id)
            nhse_spending.validate_release(nhse_spending.release_path(source), source)

    def test_top_level_finances_reconcile(self) -> None:
        nhse_spending.validate_top_level_finances()

    def test_spend_summary_matches_committed_releases_and_is_deterministic(self) -> None:
        first = nhse_spending.build_spend_summary(self.root / "first.json")
        second = nhse_spending.build_spend_summary(self.root / "second.json")
        self.assertEqual(first.read_bytes(), second.read_bytes())

        summary = json.loads(first.read_text(encoding="utf-8"))
        self.assertEqual(len(summary["years"]), len(nhse_spending.SPEND_SUMMARY_COVERAGE))
        periods = [year["period"] for year in summary["years"]]
        self.assertEqual(periods, sorted(periods))

        for year in summary["years"]:
            source = nhse_spending.load_source(year["source_id"])
            validation = json.loads(
                (nhse_spending.release_path(source) / "validation.json").read_text(encoding="utf-8")
            )
            stats = validation["statistics"]
            self.assertEqual(year["rows"], stats["rows"])
            self.assertEqual(year["total_gbp"], str(nhse_spending.Decimal(stats["total_gbp"])))
            self.assertEqual(
                year["nhs_england_total_gbp"], str(nhse_spending.Decimal(stats["nhs_england_total_gbp"]))
            )

        committed = (nhse_spending.ROOT / "spending" / "spend-by-year.json").read_bytes()
        self.assertEqual(first.read_bytes(), committed)

        full_years = [year for year in summary["years"] if year["coverage"] == "full_calendar_year"]
        partial_years = [year for year in summary["years"] if year["coverage"] != "full_calendar_year"]
        self.assertTrue(full_years)
        self.assertTrue(all(year["months_covered"] == 12 for year in full_years))
        self.assertTrue(all(year["coverage"] in {"partial_calendar_year", "fiscal_year"} for year in partial_years))

    def test_supplier_spend_summary_matches_committed_releases_and_is_deterministic(self) -> None:
        first = nhse_spending.build_supplier_spend_summary(self.root / "first-suppliers.json")
        second = nhse_spending.build_supplier_spend_summary(self.root / "second-suppliers.json")
        self.assertEqual(first.read_bytes(), second.read_bytes())

        committed = (nhse_spending.ROOT / "spending" / "supplier-spend.json").read_bytes()
        self.assertEqual(first.read_bytes(), committed)

        summary = json.loads(first.read_text(encoding="utf-8"))
        expected_periods = {
            str(nhse_spending.load_source(sid)["period"]) for sid in nhse_spending.SPEND_SUMMARY_COVERAGE
        }
        self.assertEqual({period["period"] for period in summary["periods"]}, expected_periods)
        supplier_names = [supplier["supplier"] for supplier in summary["suppliers"]]
        self.assertEqual(supplier_names, sorted(supplier_names))
        self.assertEqual(len(supplier_names), len(set(supplier_names)))

        # Every supplier's own rows/total must reconcile with its listed years, and no year
        # may reference a period outside the published coverage table.
        known_periods = {period["period"] for period in summary["periods"]}
        for supplier in summary["suppliers"]:
            self.assertEqual(supplier["total_rows"], sum(year["rows"] for year in supplier["years"]))
            self.assertEqual(
                nhse_spending.Decimal(supplier["total_gbp"]),
                sum((nhse_spending.Decimal(year["total_gbp"]) for year in supplier["years"]), nhse_spending.Decimal("0.00")),
            )
            self.assertTrue(all(year["period"] in known_periods for year in supplier["years"]))

        # Cross-check against the independently built spend-by-year.json: summing every
        # supplier's rows/total for one period must reconcile with that period's own total.
        spend_by_year = json.loads((nhse_spending.ROOT / "spending" / "spend-by-year.json").read_text(encoding="utf-8"))
        for year in spend_by_year["years"]:
            supplier_rows = sum(
                entry["rows"]
                for supplier in summary["suppliers"]
                for entry in supplier["years"]
                if entry["period"] == year["period"]
            )
            supplier_total = sum(
                (nhse_spending.Decimal(entry["total_gbp"]) for supplier in summary["suppliers"] for entry in supplier["years"] if entry["period"] == year["period"]),
                nhse_spending.Decimal("0.00"),
            )
            self.assertEqual(supplier_rows, year["rows"])
            self.assertEqual(supplier_total, nhse_spending.Decimal(year["total_gbp"]))

        # Regression check tying back to this session's manual TPP/SystmOne finding.
        phoenix = next(s for s in summary["suppliers"] if s["supplier"] == "PHOENIX PARTNERSHIP (LEEDS) LTD")
        self.assertEqual(phoenix["total_rows"], 199)
        self.assertEqual(phoenix["total_gbp"], "62839462.71")
        years_2024 = next(year for year in phoenix["years"] if year["period"] == "2024")
        self.assertEqual(years_2024["total_gbp"], "40680872.43")


class SpendingInventoryTest(unittest.TestCase):
    def test_filename_month_detection_handles_published_variants(self) -> None:
        self.assertEqual(nhse_spending_inventory._month_from_filename("25k-report-sept-2015-v2.xlsx"), 9)
        self.assertEqual(nhse_spending_inventory._month_from_filename("25k-spend-ocotober-2017.xlsx"), 10)
        self.assertIsNone(nhse_spending_inventory._month_from_filename("payments-greater-25k-report-Q4-1314.csv"))

    def test_column_profiles_ignore_physical_sheet_layout(self) -> None:
        headers = ["Entity", "Supplier", "AP Amount (£)"]
        first = {
            "format": "xlsx",
            "sheets": [
                {"name": "Payments", "header_row": 1, "headers": headers},
                {"name": "Sheet2", "header_row": None, "headers": []},
            ],
        }
        second = {
            "format": "xlsx",
            "sheets": [{"name": "Report", "header_row": 1, "headers": headers}],
        }

        self.assertNotEqual(
            nhse_spending_inventory._schema_projection(first),
            nhse_spending_inventory._schema_projection(second),
        )
        self.assertEqual(
            nhse_spending_inventory._column_projection(first),
            nhse_spending_inventory._column_projection(second),
        )

    def test_inventory_cache_is_url_specific_and_refreshable(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first = nhse_spending_inventory._cache_path(root, "https://example.test/a/report.xlsx")
            second = nhse_spending_inventory._cache_path(root, "https://example.test/b/report.xlsx")
            self.assertNotEqual(first, second)

            first.write_bytes(b"stale")
            with mock.patch.object(
                nhse_spending_inventory.urllib.request,
                "urlopen",
                return_value=io.BytesIO(b"fresh"),
            ):
                nhse_spending_inventory._download(
                    "https://example.test/a/report.xlsx",
                    first,
                    refresh=True,
                )
            self.assertEqual(first.read_bytes(), b"fresh")

    def test_committed_census_is_internally_consistent(self) -> None:
        path = Path(__file__).parent.parent / "spending" / "historical-schema-census.json"
        inventory = json.loads(path.read_text(encoding="utf-8"))
        schema_ids = {schema["id"] for schema in inventory["schemas"]}
        column_profile_ids = {profile["id"] for profile in inventory["column_profiles"]}
        source_urls = [file["source_url"] for file in inventory["files"]]

        self.assertEqual(inventory["file_count"], len(inventory["files"]))
        self.assertEqual(inventory["unique_source_count"], len(set(source_urls)))
        self.assertEqual(inventory["schema_count"], len(schema_ids))
        self.assertEqual(inventory["column_profile_count"], len(column_profile_ids))
        self.assertEqual(inventory["file_count"], 144)
        self.assertEqual(inventory["unique_source_count"], 143)
        self.assertEqual(inventory["schema_count"], 17)
        self.assertEqual(inventory["column_profile_count"], 6)
        self.assertTrue(all(file["schema_id"] in schema_ids for file in inventory["files"]))
        self.assertTrue(all(file["column_profile_id"] in column_profile_ids for file in inventory["files"]))

        schema_counts = Counter(file["schema_id"] for file in inventory["files"])
        for schema in inventory["schemas"]:
            self.assertEqual(schema["file_count"], schema_counts[schema["id"]])
            projection = nhse_spending_inventory._schema_projection(
                next(file["inspection"] for file in inventory["files"] if file["schema_id"] == schema["id"])
            )
            self.assertEqual(schema["projection"], projection)
            self.assertEqual(schema["sha256"], hashlib.sha256(nhse_spending_inventory._canonical_json(projection)).hexdigest())

        profile_counts = Counter(file["column_profile_id"] for file in inventory["files"])
        for profile in inventory["column_profiles"]:
            self.assertEqual(profile["file_count"], profile_counts[profile["id"]])
            projection = nhse_spending_inventory._column_projection(
                next(
                    file["inspection"]
                    for file in inventory["files"]
                    if file["column_profile_id"] == profile["id"]
                )
            )
            self.assertEqual(profile["projection"], projection)
            self.assertEqual(profile["sha256"], hashlib.sha256(nhse_spending_inventory._canonical_json(projection)).hexdigest())

        for collection in inventory["collections"]:
            linked = [inventory["files"][index] for index in collection["file_indexes"]]
            self.assertEqual(collection["file_count"], len(linked))
            self.assertTrue(all(file["collection_id"] == collection["id"] for file in linked))
            if collection["id"] != "2013-14":
                self.assertEqual(
                    set(collection["observed_months"]) | set(collection["missing_months"]),
                    set(range(1, 13)),
                )
                self.assertFalse(set(collection["observed_months"]) & set(collection["missing_months"]))

        duplicate_url = "https://www.england.nhs.uk/wp-content/uploads/2014/11/payments-greater-25k-report-Q4-1314.csv"
        duplicate_periods = {file["period"] for file in inventory["files"] if file["source_url"] == duplicate_url}
        self.assertEqual(duplicate_periods, {"2013-14"})


class ProcurementDataTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def _write_json(self, directory: Path, filename: str, value: dict) -> tuple[int, str]:
        path = directory / filename
        path.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        return path.stat().st_size, hashlib.sha256(path.read_bytes()).hexdigest()

    def _fixture_registry(self) -> tuple[Path, Path]:
        source_directory = self.root / "sources"
        source_directory.mkdir()
        payment_directory = self.root / "payments"
        payment_directory.mkdir()
        payment_path = payment_directory / "payments.parquet"
        pq.write_table(
            pa.table(
                {
                    "observation_id": ["payment-1", "payment-2"],
                    "entity": ["NHS England", "NHS England"],
                    "supplier": ["ACME LTD", "Unrelated supplier"],
                    "vat_registration_number": ["123456789", None],
                    "amount_gbp": [nhse_spending.Decimal("10.00"), nhse_spending.Decimal("0.00")],
                }
            ),
            payment_path,
        )
        (payment_directory / "validation.json").write_text(
            json.dumps({"statistics": {"rows": 999, "total_gbp": "999.00"}}),
            encoding="utf-8",
        )

        find_tender = {
            "extensions": [],
            "license": "https://example.test/licence",
            "packages": [],
            "publicationPolicy": "https://example.test/policy",
            "publishedDate": "2025-02-01T00:00:00Z",
            "publisher": {"name": "Fixture publisher"},
            "records": [
                {
                    "ocid": "ocds-example-1",
                    "releases": [
                        {
                            "id": "planning-1",
                            "date": "2025-01-01T00:00:00Z",
                            "tag": ["planning"],
                            "buyer": {"id": "buyer-1"},
                            "tender": {
                                "title": "Fixture procurement",
                                "value": {"amount": 100, "currency": "GBP"},
                            },
                            "parties": [
                                {
                                    "id": "buyer-1",
                                    "name": "NHS England",
                                    "identifier": {"scheme": "GB-PPON", "id": "AAAA-1111-BBBB"},
                                    "roles": ["buyer"],
                                }
                            ],
                        },
                        {
                            "id": "award-1",
                            "date": "2025-02-01T00:00:00Z",
                            "tag": ["award", "contract"],
                            "buyer": {"id": "buyer-1"},
                            "tender": {"title": "Fixture procurement"},
                            "awards": [
                                {
                                    "id": "award-a",
                                    "suppliers": [{"id": "GB-COH-00000001", "name": "ACME LIMITED"}],
                                }
                            ],
                            "contracts": [
                                {
                                    "id": "contract-a",
                                    "awardID": "award-a",
                                    "value": {"amount": 80, "currency": "GBP"},
                                }
                            ],
                            "parties": [
                                {
                                    "id": "GB-COH-00000001",
                                    "name": "ACME LIMITED",
                                    "identifier": {"scheme": "GB-COH", "id": "00000001"},
                                    "roles": ["supplier"],
                                },
                                {
                                    "id": "GB-COH-00000002",
                                    "name": "ACME LIMITED",
                                    "identifier": {"scheme": "GB-COH", "id": "00000002"},
                                    "roles": ["supplier"],
                                },
                                {
                                    "id": "GB-COH-00000003",
                                    "name": "ACME LIMITED",
                                    "identifier": {"scheme": "GB-COH", "id": "00000003"},
                                    "roles": ["tenderer"],
                                },
                            ],
                        },
                    ],
                    "compiledRelease": {},
                    "versionedRelease": {},
                }
            ],
            "uri": "https://example.test/record",
            "version": "1.1",
        }
        find_tender_bytes, find_tender_hash = self._write_json(source_directory, "find-tender.json", find_tender)

        contracts_finder = {
            "id": "replacement-guid",
            "notice": {
                "id": "replacement-guid",
                "parentId": "original-guid",
                "version": 1,
                "publishedDate": "2025-02-02T00:00:00Z",
                "title": "Fixture procurement",
                "organisationName": "NHS England",
            },
            "awards": [
                {
                    "id": "cf-award",
                    "value": 80,
                    "supplierName": "ACME LIMITED",
                    "dunsNumber": "123456789",
                }
            ],
            "additionalDetails": [],
        }
        contracts_finder_bytes, contracts_finder_hash = self._write_json(
            source_directory, "contracts-finder.json", contracts_finder
        )

        registry = {
            "schema_version": 1,
            "release": {
                "id": "fixture-procurement",
                "title": "Fixture procurement release",
                "observed_on": dt.date(2026, 10, 17),
                "normalization_schema_version": 1,
                "rights_status": "test-only",
                "rights_url": "https://example.test/licence",
                "payment_release": {
                    "id": "fixture-payments",
                    "parquet": str(payment_path),
                    "sha256": hashlib.sha256(payment_path.read_bytes()).hexdigest(),
                },
                "matching_release_ids": ["award-1"],
                "limitations": ["Test fixture"],
            },
            "sources": [
                {
                    "id": "find-tender-fixture",
                    "source_system": "find-a-tender",
                    "source_family": "find-a-tender-record",
                    "source_url": "https://example.test/record",
                    "media_type": "application/json",
                    "filename": "find-tender.json",
                    "bytes": find_tender_bytes,
                    "sha256": find_tender_hash,
                    "expected_ocid": "ocds-example-1",
                    "expected_release_ids": ["planning-1", "award-1"],
                    "monetary_observations": [
                        {
                            "id": "planning-value",
                            "release_id": "planning-1",
                            "value_type": "planning_estimate",
                            "source_pointer": "/tender/value",
                            "amount_field": "amount",
                            "expected_amount": "100.00",
                            "currency": "GBP",
                            "vat_basis": "excluding",
                            "scope_type": "procurement",
                            "scope_id": "ocds-example-1",
                        },
                        {
                            "id": "contract-value",
                            "release_id": "award-1",
                            "value_type": "contract_value",
                            "source_pointer": "/contracts/0/value",
                            "amount_field": "amount",
                            "expected_amount": "80.00",
                            "currency": "GBP",
                            "vat_basis": "excluding",
                            "scope_type": "contract",
                            "scope_id": "contract-a",
                            "commercial_event_id": "fixture-award",
                        },
                    ],
                },
                {
                    "id": "contracts-finder-fixture",
                    "source_system": "contracts-finder",
                    "source_family": "contracts-finder-published-notice-v2",
                    "source_url": "https://example.test/notice",
                    "media_type": "application/json",
                    "filename": "contracts-finder.json",
                    "bytes": contracts_finder_bytes,
                    "sha256": contracts_finder_hash,
                    "expected_notice_id": "replacement-guid",
                    "expected_parent_id": "original-guid",
                    "expected_version": 1,
                    "monetary_observations": [
                        {
                            "id": "cf-award-value",
                            "value_type": "award_value",
                            "source_pointer": "/awards/0/value",
                            "expected_amount": "80.00",
                            "currency": "GBP",
                            "vat_basis": "unspecified",
                            "scope_type": "award",
                            "scope_id": "cf-award",
                            "commercial_event_id": "fixture-award",
                        }
                    ],
                },
            ],
        }
        registry_path = self.root / "procurement-sources.yaml"
        registry_path.write_text(yaml.safe_dump(registry, sort_keys=False), encoding="utf-8")
        return registry_path, source_directory

    def test_procurement_release_build_is_deterministic(self) -> None:
        registry_path, source_directory = self._fixture_registry()
        first = nhse_procurement.build_release(self.root / "first", registry_path, source_directory)
        second = nhse_procurement.build_release(self.root / "second", registry_path, source_directory)
        self.assertEqual(
            {path.name: path.read_bytes() for path in first.iterdir()},
            {path.name: path.read_bytes() for path in second.iterdir()},
        )
        registry = nhse_procurement.load_registry(registry_path)
        nhse_procurement.validate_release(
            first,
            registry,
            registry_path=registry_path,
            source_directory=source_directory,
            verify_against_sources=True,
        )
        release = json.loads((first / "procurement.json").read_text(encoding="utf-8"))
        self.assertEqual({item["value_type"] for item in release["monetary_observations"]}, {"planning_estimate", "contract_value", "award_value"})
        self.assertTrue(all(item["is_payment"] is False for item in release["monetary_observations"]))
        self.assertEqual(len(release["match_candidates"]), 1)
        self.assertEqual(release["match_candidates"][0]["decision"], "needs_review")
        self.assertEqual(release["match_candidates"][0]["contract_attribution"], "not_asserted")
        self.assertEqual(release["canonical_payment_release"]["rows"], 2)
        self.assertEqual(release["canonical_payment_release"]["total_gbp"], "10.00")
        event_values = {
            (item["amount"], item["currency"])
            for item in release["monetary_observations"]
            if item["commercial_event_id"] == "fixture-award"
        }
        self.assertEqual(event_values, {("80.00", "GBP")})
        self.assertEqual(
            [item["relationship_type"] for item in release["relationships"]].count("replaces_notice"),
            1,
        )

    def test_procurement_release_refuses_overwrite_with_different_bytes(self) -> None:
        registry_path, source_directory = self._fixture_registry()
        target = nhse_procurement.build_release(self.root / "release", registry_path, source_directory)
        (target / "procurement.json").write_text("{}\n", encoding="utf-8")
        with self.assertRaisesRegex(nhse_procurement.ProcurementDataError, "immutable procurement release"):
            nhse_procurement.build_release(target, registry_path, source_directory)

    def test_procurement_artifact_pins_reject_coordinated_tampering(self) -> None:
        registry_path, source_directory = self._fixture_registry()
        target = nhse_procurement.build_release(self.root / "release", registry_path, source_directory)
        registry = yaml.safe_load(registry_path.read_text(encoding="utf-8"))
        registry["release"]["expected_artifacts"] = {
            filename: hashlib.sha256((target / filename).read_bytes()).hexdigest()
            for filename in ("procurement.json", "validation.json")
        }
        registry_path.write_text(yaml.safe_dump(registry, sort_keys=False), encoding="utf-8")
        registry = nhse_procurement.load_registry(registry_path)
        nhse_procurement.validate_release(target, registry, registry_path=registry_path)

        release = json.loads((target / "procurement.json").read_text(encoding="utf-8"))
        release["title"] = "Tampered title"
        (target / "procurement.json").write_text(
            json.dumps(release, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        checksums = "".join(
            f"{hashlib.sha256((target / filename).read_bytes()).hexdigest()}  {filename}\n"
            for filename in ("procurement.json", "validation.json")
        )
        (target / "checksums.sha256").write_text(checksums, encoding="ascii")
        with self.assertRaisesRegex(nhse_procurement.ProcurementDataError, "artifact digests"):
            nhse_procurement.validate_release(target, registry, registry_path=registry_path)

    def test_text_monetary_observation_is_parsed_from_source_excerpt(self) -> None:
        source = {"id": "fixture", "source_system": "find-a-tender"}
        document = {
            "records": [
                {
                    "ocid": "ocds-example-1",
                    "releases": [
                        {
                            "id": "award-1",
                            "date": "2025-01-01T00:00:00Z",
                            "tender": {"description": "The forecasted amount is £330m and is not committed."},
                        }
                    ],
                }
            ]
        }
        config = {
            "id": "forecast",
            "release_id": "award-1",
            "value_type": "forecast_not_committed",
            "source_pointer": "/tender/description",
            "source_text_contains": "forecasted amount is £330m",
            "expected_amount": "330000000.00",
            "currency": "GBP",
            "vat_basis": "unspecified",
            "scope_type": "procurement",
            "scope_id": "ocds-example-1",
        }
        observation = nhse_procurement._extract_monetary_observation(source, config, document)
        self.assertEqual(observation["amount"], "330000000.00")

        with self.assertRaisesRegex(nhse_procurement.ProcurementDataError, "amount mismatch"):
            nhse_procurement._extract_monetary_observation(
                source,
                {**config, "expected_amount": "331000000.00"},
                document,
            )

    def test_procurement_malformed_json_fails_without_raw_parser_error(self) -> None:
        registry_path, source_directory = self._fixture_registry()
        target = nhse_procurement.build_release(self.root / "release", registry_path, source_directory)
        (target / "procurement.json").write_text("{\n", encoding="utf-8")
        checksums = "".join(
            f"{hashlib.sha256((target / filename).read_bytes()).hexdigest()}  {filename}\n"
            for filename in ("procurement.json", "validation.json")
        )
        (target / "checksums.sha256").write_text(checksums, encoding="ascii")
        registry = nhse_procurement.load_registry(registry_path)
        with self.assertRaisesRegex(nhse_procurement.ProcurementDataError, "invalid procurement release JSON"):
            nhse_procurement.validate_release(target, registry, registry_path=registry_path)

    def test_procurement_malformed_structure_fails_without_key_error(self) -> None:
        registry_path, source_directory = self._fixture_registry()
        target = nhse_procurement.build_release(self.root / "release", registry_path, source_directory)
        release = json.loads((target / "procurement.json").read_text(encoding="utf-8"))
        del release["notices"][0]["notice_key"]
        (target / "procurement.json").write_text(
            json.dumps(release, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        checksums = "".join(
            f"{hashlib.sha256((target / filename).read_bytes()).hexdigest()}  {filename}\n"
            for filename in ("procurement.json", "validation.json")
        )
        (target / "checksums.sha256").write_text(checksums, encoding="ascii")
        registry = nhse_procurement.load_registry(registry_path)
        with self.assertRaisesRegex(nhse_procurement.ProcurementDataError, "malformed procurement release structure"):
            nhse_procurement.validate_release(target, registry, registry_path=registry_path)

    def test_procurement_source_hash_mismatch_fails_closed(self) -> None:
        registry_path, source_directory = self._fixture_registry()
        registry = nhse_procurement.load_registry(registry_path)
        (source_directory / "find-tender.json").write_bytes(b"{}")
        with self.assertRaisesRegex(nhse_procurement.ProcurementDataError, "byte length"):
            nhse_procurement.verify_sources(registry, source_directory)

    def test_identifier_evidence_is_scheme_aware(self) -> None:
        exact = nhse_procurement.identifier_evidence("GB-COH", "001", "GB-COH", "001")
        conflict = nhse_procurement.identifier_evidence("GB-COH", "001", "GB-COH", "002")
        cross_scheme = nhse_procurement.identifier_evidence("GB-VAT", "001", "GB-COH", "001")
        self.assertEqual(exact, {"tier": "A", "result": "same_authoritative_identifier", "contradiction": False})
        self.assertTrue(conflict["contradiction"])
        self.assertEqual(cross_scheme["tier"], "D")

    def test_committed_procurement_release_preserves_uncertainty(self) -> None:
        registry = nhse_procurement.load_registry()
        nhse_procurement.validate_release(registry=registry)
        release = json.loads((nhse_procurement.DATA_DIR / "procurement.json").read_text(encoding="utf-8"))
        candidates = release["match_candidates"]
        self.assertEqual(len(candidates), 5)
        self.assertEqual({item["supplier_text_raw"] for item in candidates}, {"PALANTIR TECHNOLOGIES UK LTD", "NOVARTIS PHARMACEUTICALS UK LTD", "PFIZER LTD"})
        self.assertTrue(all(item["decision"] == "needs_review" for item in candidates))
        self.assertTrue(all(item["evidence_tier"] == "D" for item in candidates))
        self.assertTrue(all(item["contract_attribution"] == "not_asserted" for item in candidates))
        self.assertFalse(any("NIFTY" in item["supplier_text_normalized"] for item in candidates))
        self.assertFalse(any("ACCORD" in item["notice_supplier_name_normalized"] for item in candidates))

        values = {item["monetary_observation_id"]: item for item in release["monetary_observations"]}
        self.assertEqual(values["fdp-planning-estimate"]["amount"], "240000000.00")
        self.assertEqual(values["fdp-tender-ceiling"]["amount"], "480000000.00")
        self.assertEqual(values["fdp-contract-value"]["amount"], "182242760.00")
        self.assertEqual(values["fdp-forecast-not-committed"]["amount"], "330000000.00")
        self.assertEqual(
            nhse_spending.Decimal(values["cass-modified-contract-value"]["amount"])
            - nhse_spending.Decimal(values["cass-modification-increment"]["amount"]),
            nhse_spending.Decimal("22990.00"),
        )
        self.assertTrue(all(item["is_payment"] is False for item in values.values()))
        self.assertTrue(all("amount_gbp" not in item for item in values.values()))

        replacements = [
            item for item in release["relationships"] if item["relationship_type"] == "replaces_notice"
        ]
        self.assertEqual(len(replacements), 1)
        self.assertEqual(replacements[0]["from_id"], "0f8a65b5-23a2-4294-abb1-a7fd8efb3ad0")
        self.assertEqual(replacements[0]["to_id"], "2e8c61c0-faab-4f99-ae69-b00df6bae165")


if __name__ == "__main__":
    unittest.main()
