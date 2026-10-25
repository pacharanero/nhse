#!/usr/bin/env python3
"""Fetch, normalize, and validate NHS England spending datasets."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import gzip
import hashlib
import io
import json
import os
import shutil
import sys
import tempfile
import urllib.request
from collections import Counter
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import duckdb
import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import yaml

import nhse_nhsd_spending
import nhse_procurement
from nhse_spending_inventory import InventoryError, build_inventory


ROOT = Path(__file__).resolve().parent
SOURCE_REGISTRY = ROOT / "spending" / "data-sources.yaml"
FINANCE_REGISTRY = ROOT / "spending" / "top-level-finances.yaml"
SOURCE_DIR = ROOT / "spending" / "sources"
DATA_DIR = ROOT / "spending" / "data"
TOOL_VERSION = "0.1.0"
PARQUET_ROW_GROUP_SIZE = 64_000

OUTPUT_COLUMNS_V1 = [
    "observation_id",
    "department_family",
    "entity",
    "transaction_date",
    "expense_type",
    "expense_area",
    "supplier",
    "transaction_number",
    "amount_gbp",
    "source_amount_raw",
    "currency",
    "vat_registration_number",
    "redaction_status",
    "source_sha256",
    "source_file",
    "source_sheet",
    "source_row",
]

OUTPUT_COLUMNS_V2 = [
    "observation_id",
    "department_family",
    "entity",
    "transaction_date",
    "expense_type",
    "expense_area",
    "supplier",
    "transaction_number",
    "purchase_invoice_number",
    "amount_gbp",
    "source_amount_raw",
    "currency",
    "vat_registration_number",
    "redaction_status",
    "source_id",
    "source_period",
    "source_sha256",
    "source_file",
    "source_sheet",
    "source_row",
]

REVISION_DIFF_COLUMNS = ["change_type", "row_fingerprint", "revision_source_id", *OUTPUT_COLUMNS_V2]


class SpendingDataError(RuntimeError):
    """A source, schema, or reconciliation invariant failed."""


def _jsonable(value: Any) -> Any:
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, dict):
        return {key: _jsonable(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_jsonable(item) for item in value]
    return value


def _json_bytes(value: Any) -> bytes:
    return (json.dumps(_jsonable(value), ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8")


def _canonical_json(value: Any) -> bytes:
    return json.dumps(_jsonable(value), ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode("utf-8")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _load_yaml(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as handle:
        value = yaml.safe_load(handle)
    if not isinstance(value, dict):
        raise SpendingDataError(f"{path} must contain a YAML mapping")
    return value


def load_source(source_id: str) -> dict[str, Any]:
    registry = _load_yaml(SOURCE_REGISTRY)
    if registry.get("schema_version") != 1:
        raise SpendingDataError("unsupported spending source-registry schema")
    sources = registry.get("sources")
    if not isinstance(sources, list):
        raise SpendingDataError("spending source registry must contain a sources list")
    matches = [source for source in sources if source.get("id") == source_id]
    if len(matches) != 1:
        raise SpendingDataError(f"expected exactly one source named {source_id!r}, found {len(matches)}")
    source = matches[0]
    common_required = {
        "collection_archive_url",
        "collection_url",
        "id",
        "landing_page_archive_url",
        "landing_page_url",
        "normalization_schema_version",
        "period",
        "publication_date",
        "publication_date_precision",
        "publisher",
        "referenced_licence_archive_url",
        "referenced_licence_url",
        "rights_evidence_archive_url",
        "rights_evidence_url",
        "rights_basis",
        "rights_decision_date",
        "rights_note",
        "rights_status",
        "title",
        "expected",
    }
    missing = common_required - source.keys()
    if missing:
        raise SpendingDataError(f"source {source_id!r} is missing {', '.join(sorted(missing))}")
    schema_version = int(source["normalization_schema_version"])
    if schema_version not in {1, 2}:
        raise SpendingDataError(f"source {source_id!r} uses unsupported normalization schema {schema_version}")

    workbook_required = {
        "archive_url",
        "bytes",
        "filename",
        "id",
        "period",
        "publication_date",
        "publication_date_precision",
        "sha256",
        "source_url",
        "title",
        "workbook",
        "expected",
    }
    workbooks = source.get("workbooks")
    if workbooks is None:
        missing = workbook_required - source.keys()
        if missing:
            raise SpendingDataError(f"source {source_id!r} is missing {', '.join(sorted(missing))}")
    else:
        if not isinstance(workbooks, list) or not workbooks:
            raise SpendingDataError(f"source {source_id!r} must contain at least one workbook")
        workbook_ids = [workbook.get("id") for workbook in workbooks]
        if len(workbook_ids) != len(set(workbook_ids)):
            raise SpendingDataError(f"source {source_id!r} contains duplicate workbook IDs")
        for workbook in workbooks:
            missing = workbook_required - workbook.keys()
            if missing:
                raise SpendingDataError(
                    f"workbook {workbook.get('id', '<unknown>')!r} is missing {', '.join(sorted(missing))}"
                )

    revisions = source.get("revisions", [])
    if not isinstance(revisions, list):
        raise SpendingDataError(f"source {source_id!r} revisions must be a list")
    for revision in revisions:
        missing = workbook_required - revision.keys()
        if missing:
            raise SpendingDataError(
                f"revision {revision.get('id', '<unknown>')!r} is missing {', '.join(sorted(missing))}"
            )
    current_ids = [workbook["id"] for workbook in _workbook_sources(source)]
    revision_ids = [revision["id"] for revision in revisions]
    if len(current_ids + revision_ids) != len(set(current_ids + revision_ids)):
        raise SpendingDataError(f"source {source_id!r} contains duplicate current or revision source IDs")
    revision_diff = source.get("revision_diff")
    if revision_diff is not None:
        required = {"id", "original_source_id", "corrected_source_id", "baseline_source_id"}
        if not isinstance(revision_diff, dict) or required - revision_diff.keys():
            raise SpendingDataError(f"source {source_id!r} has an incomplete revision_diff configuration")
        if revision_diff["original_source_id"] not in revision_ids:
            raise SpendingDataError(f"source {source_id!r} revision_diff refers to an unknown original source")
        for field in ("corrected_source_id", "baseline_source_id"):
            if revision_diff[field] not in current_ids:
                raise SpendingDataError(f"source {source_id!r} revision_diff {field} refers to an unknown current source")
    return source


def _output_columns(source: dict[str, Any]) -> list[str]:
    schema_version = int(source["normalization_schema_version"])
    return OUTPUT_COLUMNS_V1 if schema_version == 1 else OUTPUT_COLUMNS_V2


def _workbook_sources(source: dict[str, Any], key: str = "workbooks") -> list[dict[str, Any]]:
    workbooks = source.get(key)
    if workbooks is None:
        return [source] if key == "workbooks" else []
    shared = {
        name: source[name]
        for name in (
            "landing_page_url",
            "landing_page_archive_url",
            "collection_url",
            "collection_archive_url",
            "publisher",
            "rights_status",
            "rights_basis",
            "rights_decision_date",
            "rights_note",
            "rights_evidence_url",
            "rights_evidence_archive_url",
            "referenced_licence_url",
            "referenced_licence_archive_url",
            "normalization_schema_version",
        )
    }
    return [{**shared, **workbook} for workbook in workbooks]


def source_path(source: dict[str, Any]) -> Path:
    return SOURCE_DIR / str(source["filename"])


def release_path(source: dict[str, Any]) -> Path:
    return DATA_DIR / str(source["id"])


def verify_source_file(path: Path, source: dict[str, Any]) -> None:
    expected_bytes = int(source["bytes"])
    actual_bytes = path.stat().st_size
    if actual_bytes != expected_bytes:
        raise SpendingDataError(f"{path} is {actual_bytes} bytes; expected {expected_bytes}")
    expected_hash = str(source["sha256"])
    actual_hash = sha256_file(path)
    if actual_hash != expected_hash:
        raise SpendingDataError(f"{path} has SHA-256 {actual_hash}; expected {expected_hash}")


def fetch_source(source: dict[str, Any], use_archive: bool = False) -> Path:
    target = source_path(source)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        verify_source_file(target, source)
        return target

    url_key = "archive_url" if use_archive else "source_url"
    url = str(source[url_key])
    request = urllib.request.Request(url, headers={"User-Agent": "nhse-research-data/0.1 (+https://github.com/pacharanero/nhse)"})
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(prefix=f".{target.name}.", suffix=".part", dir=target.parent, delete=False) as output:
            temporary = Path(output.name)
            with urllib.request.urlopen(request, timeout=60) as response:
                shutil.copyfileobj(response, output, length=1024 * 1024)
            output.flush()
            os.fsync(output.fileno())
        verify_source_file(temporary, source)
        os.replace(temporary, target)
        return target
    except Exception:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        raise


def source_paths(source: dict[str, Any]) -> list[Path]:
    return [source_path(workbook) for workbook in _workbook_sources(source)]


def fetch_sources(source: dict[str, Any], use_archive: bool = False) -> list[Path]:
    paths = [fetch_source(workbook, use_archive=use_archive) for workbook in _workbook_sources(source)]
    for revision in _workbook_sources(source, "revisions"):
        fetch_source(revision, use_archive=True)
    return paths


def verify_registered_sources(source: dict[str, Any]) -> list[Path]:
    workbooks = _workbook_sources(source)
    paths = [source_path(workbook) for workbook in workbooks]
    for path, workbook in zip(paths, workbooks, strict=True):
        verify_source_file(path, workbook)
    for revision in _workbook_sources(source, "revisions"):
        verify_source_file(source_path(revision), revision)
    return paths


def _required_text(value: Any, field: str, row_number: int, permit_numeric: bool = False) -> str:
    if permit_numeric and isinstance(value, (int, float)) and not isinstance(value, bool):
        value = str(value)
    if not isinstance(value, str):
        raise SpendingDataError(f"row {row_number}: {field} must be non-empty text")
    text = value.strip()
    if not text:
        raise SpendingDataError(f"row {row_number}: {field} must be non-empty text")
    return text


def _identifier(
    value: Any, field: str, row_number: int, nullable: bool = False, permit_non_integral: bool = False
) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if nullable:
            return None
        raise SpendingDataError(f"row {row_number}: {field} must be present")
    if isinstance(value, bool):
        raise SpendingDataError(f"row {row_number}: {field} cannot be boolean")
    if isinstance(value, float):
        if not value.is_integer():
            if permit_non_integral:
                return str(value)
            raise SpendingDataError(f"row {row_number}: {field} has a non-integral numeric identifier")
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _date(value: Any, row_number: int) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    raise SpendingDataError(f"row {row_number}: Date must be an Excel date, found {type(value).__name__}")


def _money(value: Any, row_number: int, scale: int) -> tuple[Decimal, str]:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise SpendingDataError(f"row {row_number}: AP Amount (£) must be numeric")
    if scale != 2:
        raise SpendingDataError(f"row {row_number}: only two-decimal GBP sources are currently supported")
    try:
        source_amount = Decimal(str(value))
        amount = source_amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except InvalidOperation as error:
        raise SpendingDataError(f"row {row_number}: AP Amount (£) is not finite") from error
    if not amount.is_finite():
        raise SpendingDataError(f"row {row_number}: AP Amount (£) is not finite")
    return amount, format(source_amount, "f")


def parse_workbook(path: Path, source: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    verify_source_file(path, source)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheet_name = str(source["workbook"]["sheet"])
        if sheet_name not in workbook.sheetnames:
            raise SpendingDataError(f"workbook does not contain expected sheet {sheet_name!r}")
        sheet = workbook[sheet_name]
        expected_headers = list(source["workbook"]["columns"])
        amount_scale = int(source["workbook"]["amount_scale"])
        ignore_trailing_columns = bool(source["workbook"].get("ignore_trailing_columns", False))
        excluded_rows = {int(value) for value in source["workbook"].get("excluded_rows", [])}
        if excluded_rows and not source["workbook"].get("excluded_rows_reason"):
            raise SpendingDataError("excluded_rows requires an excluded_rows_reason")
        if sheet.max_column < len(expected_headers):
            raise SpendingDataError(f"sheet has {sheet.max_column} columns; expected at least {len(expected_headers)}")
        if sheet.max_column > len(expected_headers) and not ignore_trailing_columns:
            raise SpendingDataError(
                f"sheet has {sheet.max_column} columns; expected {len(expected_headers)}. "
                "Set workbook.ignore_trailing_columns only after inspecting what the extra columns contain."
            )

        rows = sheet.iter_rows(min_row=1, max_col=len(expected_headers))
        header_cells = next(rows, None)
        if header_cells is None:
            raise SpendingDataError("workbook sheet is empty")
        actual_headers = [cell.value for cell in header_cells]
        if actual_headers != expected_headers:
            raise SpendingDataError(f"unexpected workbook headers: {actual_headers!r}")

        records: list[dict[str, Any]] = []
        excluded_count = 0
        for row_number, cells in enumerate(rows, start=2):
            if row_number in excluded_rows:
                excluded_count += 1
                continue
            values = [cell.value for cell in cells]
            if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
                continue
            formula_columns = [expected_headers[index] for index, cell in enumerate(cells) if cell.data_type == "f"]
            if formula_columns:
                raise SpendingDataError(f"row {row_number}: formulas are not allowed in {', '.join(formula_columns)}")
            error_columns = [expected_headers[index] for index, cell in enumerate(cells) if cell.data_type == "e"]
            if error_columns:
                raise SpendingDataError(f"row {row_number}: Excel errors are not allowed in {', '.join(error_columns)}")

            permit_numeric_text_fields = bool(source["workbook"].get("permit_numeric_text_fields", False))
            department_family = _required_text(values[0], "Department family", row_number, permit_numeric_text_fields)
            entity = _required_text(values[1], "Entity", row_number, permit_numeric_text_fields)
            transaction_date = _date(values[2], row_number)
            expense_type = _required_text(values[3], "Expense Type for Report", row_number, permit_numeric_text_fields)
            expense_area = _required_text(values[4], "Expense Area for Report", row_number, permit_numeric_text_fields)
            supplier = _required_text(values[5], "Supplier", row_number, permit_numeric_text_fields)
            schema_version = int(source["normalization_schema_version"])
            no_transaction_number = bool(source["workbook"].get("no_transaction_number_column", False))
            no_vat_registration_number = bool(source["workbook"].get("no_vat_registration_number_column", False))
            if no_transaction_number and no_vat_registration_number:
                raise SpendingDataError("no_transaction_number_column and no_vat_registration_number_column are mutually exclusive")
            if (no_transaction_number or no_vat_registration_number) and schema_version < 2:
                raise SpendingDataError(
                    "no_transaction_number_column and no_vat_registration_number_column require normalization schema v2"
                )
            permit_non_integral = bool(source["workbook"].get("permit_non_integral_purchase_invoice_numbers", False))
            purchase_invoice_number = None
            if no_transaction_number:
                transaction_number = None
                amount, source_amount_raw = _money(values[6], row_number, amount_scale)
                vat_number = _identifier(values[7], "VAT registration number", row_number, nullable=True)
                purchase_invoice_number = _identifier(
                    values[8], "Purchase invoice number", row_number, nullable=True, permit_non_integral=permit_non_integral
                )
            elif no_vat_registration_number:
                transaction_number = _identifier(values[6], "Transaction number", row_number)
                amount, source_amount_raw = _money(values[7], row_number, amount_scale)
                vat_number = None
                purchase_invoice_number = _identifier(
                    values[8], "Purchase invoice number", row_number, nullable=True, permit_non_integral=permit_non_integral
                )
            else:
                transaction_number = _identifier(values[6], "Transaction number", row_number)
                amount, source_amount_raw = _money(values[7], row_number, amount_scale)
                vat_number = _identifier(values[8], "VAT registration number", row_number, nullable=True)
                if schema_version >= 2:
                    if len(values) != 10:
                        raise SpendingDataError(f"row {row_number}: schema v2 requires a purchase invoice number column")
                    purchase_invoice_number = _identifier(
                        values[9],
                        "Purchase invoice number",
                        row_number,
                        nullable=True,
                        permit_non_integral=permit_non_integral,
                    )
            source_values = {
                "department_family": department_family,
                "entity": entity,
                "transaction_date": transaction_date,
                "expense_type": expense_type,
                "expense_area": expense_area,
                "supplier": supplier,
                "transaction_number": transaction_number,
                "amount_gbp": amount,
                "source_amount_raw": source_amount_raw,
                "vat_registration_number": vat_number,
            }
            if schema_version >= 2:
                source_values["purchase_invoice_number"] = purchase_invoice_number
            identity_input = {
                "schema_version": schema_version,
                "source_sha256": source["sha256"],
                "source_sheet": sheet_name,
                "source_row": row_number,
                "values": source_values,
            }
            redacted = any(isinstance(value, str) and value.strip().casefold() == "removed" for value in values)
            record = {
                "observation_id": hashlib.sha256(_canonical_json(identity_input)).hexdigest(),
                **source_values,
                "currency": "GBP",
                "redaction_status": "source_removed" if redacted else "none",
                "source_sha256": str(source["sha256"]),
                "source_file": str(source["filename"]),
                "source_sheet": sheet_name,
                "source_row": row_number,
            }
            if schema_version >= 2:
                record["source_id"] = str(source["id"])
                record["source_period"] = str(source["period"])
            records.append(record)
    finally:
        workbook.close()

    records.sort(
        key=lambda row: (
            row.get("source_period", row["source_sha256"]),
            row["source_sheet"],
            row["source_row"],
            row["observation_id"],
        )
    )
    if len({row["observation_id"] for row in records}) != len(records):
        raise SpendingDataError("observation IDs are not unique")
    if excluded_count != len(excluded_rows):
        raise SpendingDataError(
            f"registered excluded_rows did not match {excluded_count} encountered rows against "
            f"{len(excluded_rows)} registered row numbers"
        )

    entity_counts = Counter(row["entity"] for row in records)
    entity_totals: Counter[str] = Counter()
    for row in records:
        entity_totals[row["entity"]] += row["amount_gbp"]
    stats = {
        "rows": len(records),
        "total_gbp": sum((row["amount_gbp"] for row in records), Decimal("0.00")),
        "nhs_england_rows": entity_counts["NHS England"],
        "nhs_england_total_gbp": entity_totals["NHS England"],
        "negative_rows": sum(row["amount_gbp"] < 0 for row in records),
        "missing_vat_rows": sum(row["vat_registration_number"] is None for row in records),
        "minimum_date": min(row["transaction_date"] for row in records),
        "maximum_date": max(row["transaction_date"] for row in records),
        "entity_counts": dict(sorted(entity_counts.items())),
        "entity_totals_gbp": dict(sorted(entity_totals.items())),
    }
    if excluded_count:
        stats["excluded_malformed_rows"] = excluded_count
    _validate_expected_stats(stats, source["expected"])
    return records, stats


def parse_release(
    paths: list[Path], source: dict[str, Any]
) -> tuple[list[dict[str, Any]], dict[str, Any], dict[str, dict[str, Any]]]:
    workbooks = _workbook_sources(source)
    if len(paths) != len(workbooks):
        raise SpendingDataError(f"received {len(paths)} source paths for {len(workbooks)} registered workbooks")

    records: list[dict[str, Any]] = []
    source_statistics: dict[str, dict[str, Any]] = {}
    for path, workbook in zip(paths, workbooks, strict=True):
        workbook_records, workbook_stats = parse_workbook(path, workbook)
        records.extend(workbook_records)
        source_statistics[str(workbook["id"])] = workbook_stats

    records.sort(
        key=lambda row: (
            row.get("source_period", row["source_sha256"]),
            row["source_sheet"],
            row["source_row"],
            row["observation_id"],
        )
    )
    if len({row["observation_id"] for row in records}) != len(records):
        raise SpendingDataError("observation IDs are not unique across the release")
    stats = _record_statistics(records)
    _validate_expected_stats(stats, source["expected"])
    return records, stats, source_statistics


def _record_statistics(records: list[dict[str, Any]]) -> dict[str, Any]:
    entity_counts = Counter(row["entity"] for row in records)
    entity_totals: Counter[str] = Counter()
    for row in records:
        entity_totals[row["entity"]] += row["amount_gbp"]
    return {
        "rows": len(records),
        "total_gbp": sum((row["amount_gbp"] for row in records), Decimal("0.00")),
        "nhs_england_rows": entity_counts["NHS England"],
        "nhs_england_total_gbp": entity_totals["NHS England"],
        "negative_rows": sum(row["amount_gbp"] < 0 for row in records),
        "missing_vat_rows": sum(row["vat_registration_number"] is None for row in records),
        "minimum_date": min(row["transaction_date"] for row in records),
        "maximum_date": max(row["transaction_date"] for row in records),
        "entity_counts": dict(sorted(entity_counts.items())),
        "entity_totals_gbp": dict(sorted(entity_totals.items())),
    }


def _validate_expected_stats(stats: dict[str, Any], expected: dict[str, Any]) -> None:
    decimal_fields = {"total_gbp", "nhs_england_total_gbp"}
    for field, expected_value in expected.items():
        if field not in stats:
            raise SpendingDataError(f"{field} is not a recognized statistic for this workbook's rows")
        actual_value = stats[field]
        if field in decimal_fields:
            expected_value = Decimal(str(expected_value))
        else:
            expected_value = int(expected_value)
        if actual_value != expected_value:
            raise SpendingDataError(f"{field} is {actual_value}; expected {expected_value}")


def _arrow_schema(
    provenance: dict[str, Any], columns: list[str], transaction_number_nullable: bool = False
) -> pa.Schema:
    types: dict[str, tuple[pa.DataType, bool]] = {
        "observation_id": (pa.string(), False),
        "department_family": (pa.string(), False),
        "entity": (pa.string(), False),
        "transaction_date": (pa.date32(), False),
        "expense_type": (pa.string(), False),
        "expense_area": (pa.string(), False),
        "supplier": (pa.string(), False),
        "transaction_number": (pa.string(), transaction_number_nullable),
        "purchase_invoice_number": (pa.string(), True),
        "amount_gbp": (pa.decimal128(18, 2), False),
        "source_amount_raw": (pa.string(), False),
        "currency": (pa.string(), False),
        "vat_registration_number": (pa.string(), True),
        "redaction_status": (pa.string(), False),
        "source_id": (pa.string(), False),
        "source_period": (pa.string(), False),
        "source_sha256": (pa.string(), False),
        "source_file": (pa.string(), False),
        "source_sheet": (pa.string(), False),
        "source_row": (pa.int32(), False),
    }
    fields = [pa.field(column, types[column][0], nullable=types[column][1]) for column in columns]
    return pa.schema(fields, metadata={b"nhse.provenance": _canonical_json(provenance)})


def _write_csv(path: Path, records: list[dict[str, Any]], columns: list[str]) -> None:
    with path.open("wb") as raw_output:
        with gzip.GzipFile(filename="", mode="wb", fileobj=raw_output, mtime=0, compresslevel=9) as compressed:
            with io.TextIOWrapper(compressed, encoding="utf-8", newline="") as text_output:
                writer = csv.DictWriter(text_output, fieldnames=columns, lineterminator="\n")
                writer.writeheader()
                for record in records:
                    writer.writerow({column: _jsonable(record[column]) for column in columns})


def _content_fingerprint(records: list[dict[str, Any]]) -> str:
    digest = hashlib.sha256()
    for record in records:
        digest.update(_canonical_json(record))
        digest.update(b"\n")
    return digest.hexdigest()


def _typed_csv_record(record: dict[str, str]) -> dict[str, Any]:
    typed: dict[str, Any] = dict(record)
    if "transaction_date" in typed:
        typed["transaction_date"] = dt.date.fromisoformat(record["transaction_date"])
    if "amount_gbp" in typed:
        typed["amount_gbp"] = Decimal(record["amount_gbp"])
    if "source_row" in typed:
        typed["source_row"] = int(record["source_row"])
    for field in ("purchase_invoice_number", "vat_registration_number", "transaction_number"):
        if field in typed:
            typed[field] = record[field] or None
    return typed


def _expected_observation_id(record: dict[str, Any], schema_version: int) -> str:
    value_fields = [
        "department_family",
        "entity",
        "transaction_date",
        "expense_type",
        "expense_area",
        "supplier",
        "transaction_number",
        "amount_gbp",
        "source_amount_raw",
        "vat_registration_number",
    ]
    if schema_version >= 2:
        value_fields.append("purchase_invoice_number")
    identity_input = {
        "schema_version": schema_version,
        "source_sha256": record["source_sha256"],
        "source_sheet": record["source_sheet"],
        "source_row": record["source_row"],
        "values": {field: record[field] for field in value_fields},
    }
    return hashlib.sha256(_canonical_json(identity_input)).hexdigest()


def _resource_schema(columns: list[str], primary_key: str | list[str] = "observation_id") -> dict[str, Any]:
    types = {
        "change_type": "string",
        "row_fingerprint": "string",
        "revision_source_id": "string",
        "observation_id": "string",
        "department_family": "string",
        "entity": "string",
        "transaction_date": "date",
        "expense_type": "string",
        "expense_area": "string",
        "supplier": "string",
        "transaction_number": "string",
        "purchase_invoice_number": "string",
        "amount_gbp": "number",
        "source_amount_raw": "string",
        "currency": "string",
        "vat_registration_number": "string",
        "redaction_status": "string",
        "source_id": "string",
        "source_period": "string",
        "source_sha256": "string",
        "source_file": "string",
        "source_sheet": "string",
        "source_row": "integer",
    }
    return {
        "fields": [
            {"name": column, "type": types[column], **({"decimalChar": "."} if column == "amount_gbp" else {})}
            for column in columns
        ],
        "primaryKey": primary_key,
    }


def _source_provenance(source: dict[str, Any]) -> dict[str, Any]:
    return {
        key: _jsonable(source[key])
        for key in (
            "id",
            "title",
            "period",
            "publication_date",
            "publication_date_precision",
            "publisher",
            "landing_page_url",
            "landing_page_archive_url",
            "collection_url",
            "collection_archive_url",
            "source_url",
            "archive_url",
            "filename",
            "sha256",
            "bytes",
            "rights_status",
            "rights_basis",
            "rights_decision_date",
            "rights_note",
            "rights_evidence_url",
            "rights_evidence_archive_url",
            "referenced_licence_url",
            "referenced_licence_archive_url",
        )
    }


def _semantic_row(record: dict[str, Any], *, include_vat: bool = True) -> dict[str, Any]:
    fields = [
        "department_family",
        "entity",
        "transaction_date",
        "expense_type",
        "expense_area",
        "supplier",
        "transaction_number",
        "purchase_invoice_number",
        "amount_gbp",
        "source_amount_raw",
        "currency",
        "redaction_status",
    ]
    if include_vat:
        fields.append("vat_registration_number")
    return {field: record[field] for field in fields}


def _revision_diff(
    source: dict[str, Any], corrected_records: list[dict[str, Any]], current_records: list[dict[str, Any]]
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    config = source["revision_diff"]
    revisions = {revision["id"]: revision for revision in _workbook_sources(source, "revisions")}
    workbooks = {workbook["id"]: workbook for workbook in _workbook_sources(source)}
    try:
        original = revisions[str(config["original_source_id"])]
        corrected = workbooks[str(config["corrected_source_id"])]
        baseline = workbooks[str(config["baseline_source_id"])]
    except (KeyError, TypeError) as error:
        raise SpendingDataError("revision_diff refers to an unknown or invalid source") from error

    original_records, original_stats = parse_workbook(source_path(original), original)
    corrected_rows = [row for row in corrected_records if row["source_id"] == corrected["id"]]
    baseline_rows = [row for row in current_records if row["source_id"] == baseline["id"]]

    def indexed(records: list[dict[str, Any]]) -> dict[bytes, list[dict[str, Any]]]:
        result: dict[bytes, list[dict[str, Any]]] = {}
        for record in records:
            result.setdefault(_canonical_json(_semantic_row(record)), []).append(record)
        return result

    original_index = indexed(original_records)
    corrected_index = indexed(corrected_rows)
    semantic_keys = sorted(original_index.keys() | corrected_index.keys())
    diff_rows: list[dict[str, Any]] = []
    unchanged_rows = 0
    for semantic_key in semantic_keys:
        before = original_index.get(semantic_key, [])
        after = corrected_index.get(semantic_key, [])
        unchanged = min(len(before), len(after))
        unchanged_rows += unchanged
        for change_type, revision_source_id, records in (
            ("removed", original["id"], before[unchanged:]),
            ("added", corrected["id"], after[unchanged:]),
        ):
            for record in records:
                diff_rows.append(
                    {
                        "change_type": change_type,
                        "row_fingerprint": hashlib.sha256(semantic_key).hexdigest(),
                        "revision_source_id": revision_source_id,
                        **record,
                    }
                )
    diff_rows.sort(
        key=lambda row: (row["change_type"], row["row_fingerprint"], row["source_row"], row["observation_id"])
    )

    original_without_vat = [_canonical_json(_semantic_row(row, include_vat=False)) for row in original_records]
    baseline_without_vat = [_canonical_json(_semantic_row(row, include_vat=False)) for row in baseline_rows]
    candidate_fields = ["entity", "transaction_date", "transaction_number", "supplier", "purchase_invoice_number"]

    def candidate_keys(records: list[dict[str, Any]]) -> Counter[bytes]:
        return Counter(_canonical_json({field: row[field] for field in candidate_fields}) for row in records)

    shared_candidate_keys = sum((candidate_keys(original_records) & candidate_keys(corrected_rows)).values())
    summary = {
        "schema_version": 1,
        "status": "complete",
        "revision_id": str(config["id"]),
        "original_source": _source_provenance(original),
        "corrected_source": _source_provenance(corrected),
        "comparison": {
            "method": "exact multiset comparison of normalized analytical rows",
            "candidate_key_fields": candidate_fields,
            "original_rows": len(original_records),
            "corrected_rows": len(corrected_rows),
            "unchanged_rows": unchanged_rows,
            "removed_rows": sum(row["change_type"] == "removed" for row in diff_rows),
            "added_rows": sum(row["change_type"] == "added" for row in diff_rows),
            "shared_candidate_keys": shared_candidate_keys,
        },
        "original_statistics": _jsonable(original_stats),
        "corrected_statistics": _jsonable(_record_statistics(corrected_rows)),
        "diagnosis": {
            "baseline_source_id": baseline["id"],
            "ordered_rows_equal_to_baseline_excluding_vat": original_without_vat == baseline_without_vat,
            "baseline_rows": len(baseline_rows),
            "original_missing_vat_rows": original_stats["missing_vat_rows"],
            "baseline_missing_vat_rows": _record_statistics(baseline_rows)["missing_vat_rows"],
            "note": "The superseded July workbook repeats the June rows apart from VAT identifiers; the corrected workbook replaces them with July-dated rows.",
        },
    }
    return diff_rows, summary


def build_release(path: Path | list[Path], source: dict[str, Any], target: Path) -> Path:
    paths = [path] if isinstance(path, Path) else path
    records, stats, source_statistics = parse_release(paths, source)
    columns = _output_columns(source)
    schema_version = int(source["normalization_schema_version"])
    content_fingerprint = _content_fingerprint(records)
    provenance = {
        "dataset": "nhs-england-payments-over-25000",
        "dataset_version": f"{source['period']}.1",
        "schema_version": schema_version,
        "tool": "nhse-research-data",
        "tool_version": TOOL_VERSION,
        "normalized_content_sha256": content_fingerprint,
    }
    if "workbooks" in source:
        provenance["sources"] = [_source_provenance(workbook) for workbook in _workbook_sources(source)]
    else:
        provenance["source"] = _source_provenance(source)

    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(tempfile.mkdtemp(prefix=f".{target.name}.", dir=target.parent))
    backup = target.with_name(f".{target.name}.previous")
    try:
        parquet_path = temporary / "payments.parquet"
        csv_path = temporary / "payments.csv.gz"
        transaction_number_nullable = any(record.get("transaction_number") is None for record in records)
        table = pa.Table.from_pylist(
            records, schema=_arrow_schema(provenance, columns, transaction_number_nullable)
        )
        pq.write_table(
            table,
            parquet_path,
            compression="zstd",
            compression_level=3,
            row_group_size=PARQUET_ROW_GROUP_SIZE,
            version="2.6",
            write_statistics=True,
        )
        _write_csv(csv_path, records, columns)

        parquet_hash = sha256_file(parquet_path)
        csv_hash = sha256_file(csv_path)
        validation = {
            "dataset": "nhs-england-payments-over-25000",
            "dataset_version": f"{source['period']}.1",
            "schema_version": schema_version,
            "normalized_content_sha256": content_fingerprint,
            "artifacts": {
                "payments.csv.gz": {"bytes": csv_path.stat().st_size, "sha256": csv_hash},
                "payments.parquet": {"bytes": parquet_path.stat().st_size, "sha256": parquet_hash},
            },
            "statistics": _jsonable(stats),
            "status": "passed",
        }
        if "workbooks" in source:
            validation["source_statistics"] = _jsonable(source_statistics)
            validation["source_sha256s"] = [workbook["sha256"] for workbook in _workbook_sources(source)]
        else:
            validation["source_sha256"] = source["sha256"]
        descriptor = {
            "profile": "tabular-data-package",
            "name": f"nhs-england-payments-over-25000-{source['period']}",
            "title": source["title"],
            "version": f"{source['period']}.1",
            "description": "Normalized NHS England payments-over-£25,000 transparency data with physical workbook-row provenance.",
            "licenses": [
                {
                    "name": "OGL-UK-3.0",
                    "path": source["referenced_licence_url"],
                    "title": "Open Government Licence v3.0",
                }
            ],
            "rights": source["rights_note"],
            "rights_status": source["rights_status"],
            "sources": [_source_provenance(workbook) for workbook in _workbook_sources(source)],
            "nhse": {
                "normalized_content_sha256": content_fingerprint,
                "schema_version": schema_version,
                "tool_version": TOOL_VERSION,
                "limitations": [
                    "This thresholded payments publication is not a general ledger and does not reconcile one-for-one to annual accrual accounts.",
                    "Supplier text is preserved as published and has not been entity-resolved.",
                    "Negative values are retained and not classified as refunds, reversals or corrections without source evidence.",
                    "The project-owner redistribution decision covers this normalized release, not a raw-workbook mirror.",
                ],
            },
            "resources": [
                {
                    "name": "payments-parquet",
                    "path": "payments.parquet",
                    "format": "parquet",
                    "mediatype": "application/vnd.apache.parquet",
                    "bytes": parquet_path.stat().st_size,
                    "hash": f"sha256:{parquet_hash}",
                    "schema": _resource_schema(columns),
                },
                {
                    "name": "payments-csv",
                    "path": "payments.csv.gz",
                    "format": "csv",
                    "mediatype": "text/csv",
                    "compression": "gzip",
                    "encoding": "utf-8",
                    "bytes": csv_path.stat().st_size,
                    "hash": f"sha256:{csv_hash}",
                    "schema": _resource_schema(columns),
                },
            ],
        }

        if "revision_diff" in source:
            diff_rows, revision_summary = _revision_diff(source, records, records)
            revision_path = temporary / "july-2024-correction.json"
            diff_path = temporary / "july-2024-correction-diff.csv.gz"
            revision_path.write_bytes(_json_bytes(revision_summary))
            _write_csv(diff_path, diff_rows, REVISION_DIFF_COLUMNS)
            validation["revision"] = revision_summary["comparison"]
            validation["artifacts"][diff_path.name] = {
                "bytes": diff_path.stat().st_size,
                "sha256": sha256_file(diff_path),
            }
            validation["artifacts"][revision_path.name] = {
                "bytes": revision_path.stat().st_size,
                "sha256": sha256_file(revision_path),
            }
            descriptor["nhse"]["revision"] = revision_summary
            descriptor["resources"].extend(
                [
                    {
                        "name": "july-2024-correction-diff",
                        "path": diff_path.name,
                        "format": "csv",
                        "mediatype": "text/csv",
                        "compression": "gzip",
                        "encoding": "utf-8",
                        "bytes": diff_path.stat().st_size,
                        "hash": f"sha256:{sha256_file(diff_path)}",
                        "schema": _resource_schema(
                            REVISION_DIFF_COLUMNS,
                            primary_key=["change_type", "observation_id"],
                        ),
                    },
                    {
                        "name": "july-2024-correction-summary",
                        "path": revision_path.name,
                        "format": "json",
                        "mediatype": "application/json",
                        "bytes": revision_path.stat().st_size,
                        "hash": f"sha256:{sha256_file(revision_path)}",
                    },
                ]
            )
        (temporary / "validation.json").write_bytes(_json_bytes(validation))
        (temporary / "datapackage.json").write_bytes(_json_bytes(descriptor))

        checksummed = ["datapackage.json", "payments.csv.gz", "payments.parquet", "validation.json"]
        if "revision_diff" in source:
            checksummed.extend(["july-2024-correction-diff.csv.gz", "july-2024-correction.json"])
        checksum_lines = [f"{sha256_file(temporary / name)}  {name}\n" for name in sorted(checksummed)]
        (temporary / "checksums.sha256").write_text("".join(checksum_lines), encoding="ascii")
        validate_release(temporary, source, verify_against_sources=True, source_files=paths)

        if backup.exists():
            shutil.rmtree(backup)
        if target.exists():
            target.rename(backup)
        temporary.rename(target)
        if backup.exists():
            shutil.rmtree(backup)
        return target
    except Exception:
        if temporary.exists():
            shutil.rmtree(temporary)
        if backup.exists() and not target.exists():
            backup.rename(target)
        raise


def validate_release(
    target: Path,
    source: dict[str, Any],
    *,
    verify_against_sources: bool = False,
    source_files: list[Path] | None = None,
) -> None:
    required_files = {"checksums.sha256", "datapackage.json", "payments.csv.gz", "payments.parquet", "validation.json"}
    if "revision_diff" in source:
        required_files.update({"july-2024-correction-diff.csv.gz", "july-2024-correction.json"})
    actual_files = {path.name for path in target.iterdir() if path.is_file()}
    if actual_files != required_files:
        raise SpendingDataError(f"release files are {sorted(actual_files)}; expected {sorted(required_files)}")

    checksum_lines = (target / "checksums.sha256").read_text(encoding="ascii").splitlines()
    expected_checksums = required_files - {"checksums.sha256"}
    checksummed_names = [line.split("  ", 1)[1] for line in checksum_lines]
    if len(checksummed_names) != len(set(checksummed_names)) or set(checksummed_names) != expected_checksums:
        raise SpendingDataError(f"checksums cover {sorted(checksummed_names)}; expected {sorted(expected_checksums)}")
    for line in checksum_lines:
        digest, filename = line.split("  ", 1)
        if sha256_file(target / filename) != digest:
            raise SpendingDataError(f"checksum mismatch for {filename}")

    descriptor = json.loads((target / "datapackage.json").read_text(encoding="utf-8"))
    if descriptor.get("rights_status") != source["rights_status"]:
        raise SpendingDataError("Data Package rights status does not match the source registry")
    expected_sources = [_source_provenance(workbook) for workbook in _workbook_sources(source)]
    if descriptor.get("sources") != expected_sources:
        raise SpendingDataError("Data Package sources do not match the source registry")

    validation = json.loads((target / "validation.json").read_text(encoding="utf-8"))
    if "workbooks" in source:
        expected_source_sha256s = [workbook["sha256"] for workbook in _workbook_sources(source)]
        if validation.get("source_sha256s") != expected_source_sha256s:
            raise SpendingDataError("validation source digests do not match the source registry")
    elif validation.get("source_sha256") != source["sha256"]:
        raise SpendingDataError("validation source digest does not match the source registry")
    parquet_records = pq.read_table(target / "payments.parquet").to_pylist()
    content_fingerprint = _content_fingerprint(parquet_records)
    if validation.get("normalized_content_sha256") != content_fingerprint:
        raise SpendingDataError("validation content fingerprint does not match Parquet rows")
    if descriptor.get("nhse", {}).get("normalized_content_sha256") != content_fingerprint:
        raise SpendingDataError("Data Package content fingerprint does not match Parquet rows")
    observation_ids = [record["observation_id"] for record in parquet_records]
    if len(observation_ids) != len(set(observation_ids)):
        raise SpendingDataError("Parquet observation IDs are not unique")
    if any(record["source_row"] < 1 for record in parquet_records):
        raise SpendingDataError("Parquet contains an invalid source row")
    if any(record["currency"] != "GBP" for record in parquet_records):
        raise SpendingDataError("Parquet contains a non-GBP currency")
    if any(record["redaction_status"] not in {"none", "source_removed"} for record in parquet_records):
        raise SpendingDataError("Parquet contains an invalid redaction status")

    schema_version = int(source["normalization_schema_version"])
    if any(record["observation_id"] != _expected_observation_id(record, schema_version) for record in parquet_records):
        raise SpendingDataError("Parquet observation ID does not match its values and source provenance")
    if schema_version == 2:
        expected_bindings = {
            (
                workbook["id"],
                str(workbook["period"]),
                workbook["sha256"],
                workbook["filename"],
                workbook["workbook"]["sheet"],
            )
            for workbook in _workbook_sources(source)
        }
        actual_bindings = {
            (
                record["source_id"],
                record["source_period"],
                record["source_sha256"],
                record["source_file"],
                record["source_sheet"],
            )
            for record in parquet_records
        }
    else:
        expected_bindings = {
            (source["sha256"], source["filename"], source["workbook"]["sheet"]),
        }
        actual_bindings = {
            (record["source_sha256"], record["source_file"], record["source_sheet"])
            for record in parquet_records
        }
    if actual_bindings != expected_bindings:
        raise SpendingDataError("Parquet row provenance does not match the source registry")
    if schema_version == 2:
        for workbook in _workbook_sources(source):
            workbook_records = [record for record in parquet_records if record["source_id"] == workbook["id"]]
            # excluded_malformed_rows describes rows skipped before the Parquet was written, so it
            # cannot be rederived from committed rows; parse_workbook already checks it at parse time.
            workbook_expected = {
                key: value for key, value in workbook["expected"].items() if key != "excluded_malformed_rows"
            }
            _validate_expected_stats(_record_statistics(workbook_records), workbook_expected)
    if verify_against_sources:
        expected_records, _, _ = parse_release(source_files or source_paths(source), source)
        if parquet_records != expected_records:
            raise SpendingDataError("Parquet rows do not match the registered source workbooks")

    connection = duckdb.connect(":memory:")
    try:
        row = connection.execute(
            """
            SELECT count(*), sum(amount_gbp),
                   count(*) FILTER (WHERE entity = 'NHS England'),
                   sum(amount_gbp) FILTER (WHERE entity = 'NHS England'),
                   count(*) FILTER (WHERE amount_gbp < 0),
                   count(*) FILTER (WHERE vat_registration_number IS NULL)
            FROM read_parquet(?)
            """,
            [str(target / "payments.parquet")],
        ).fetchone()
    finally:
        connection.close()
    parquet_stats = {
        "rows": row[0],
        "total_gbp": row[1],
        "nhs_england_rows": row[2],
        "nhs_england_total_gbp": row[3],
        "negative_rows": row[4],
        "missing_vat_rows": row[5],
    }
    _validate_expected_stats(parquet_stats, source["expected"])

    csv_rows = 0
    csv_total = Decimal("0.00")
    csv_records: list[dict[str, Any]] = []
    with gzip.open(target / "payments.csv.gz", "rt", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != _output_columns(source):
            raise SpendingDataError("CSV column order does not match the canonical schema")
        for record in reader:
            csv_rows += 1
            csv_total += Decimal(record["amount_gbp"])
            csv_records.append(_typed_csv_record(record))
    if csv_rows != parquet_stats["rows"] or csv_total != parquet_stats["total_gbp"]:
        raise SpendingDataError("CSV and Parquet row counts or totals differ")
    if csv_records != parquet_records:
        raise SpendingDataError("CSV and Parquet normalized rows or row order differ")

    if "revision_diff" in source:
        revision = json.loads((target / "july-2024-correction.json").read_text(encoding="utf-8"))
        comparison = revision["comparison"]
        config = source["revision_diff"]
        if revision.get("revision_id") != str(config["id"]):
            raise SpendingDataError("revision summary ID does not match the source registry")
        if revision.get("original_source", {}).get("id") != str(config["original_source_id"]):
            raise SpendingDataError("revision summary original source does not match the source registry")
        if revision.get("corrected_source", {}).get("id") != str(config["corrected_source_id"]):
            raise SpendingDataError("revision summary corrected source does not match the source registry")
        resources = {resource["name"]: resource for resource in descriptor.get("resources", [])}
        expected_diff_schema = _resource_schema(
            REVISION_DIFF_COLUMNS,
            primary_key=["change_type", "observation_id"],
        )
        if resources.get("july-2024-correction-diff", {}).get("schema") != expected_diff_schema:
            raise SpendingDataError("revision-diff Data Package schema does not match its CSV schema")
        diff_counts: Counter[str] = Counter()
        typed_diff_rows: list[dict[str, Any]] = []
        with gzip.open(target / "july-2024-correction-diff.csv.gz", "rt", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames != REVISION_DIFF_COLUMNS:
                raise SpendingDataError("revision-diff CSV column order does not match its schema")
            for record in reader:
                change_type = record["change_type"]
                if change_type not in {"added", "removed"}:
                    raise SpendingDataError(f"unknown revision change type {change_type!r}")
                expected_source_id = str(
                    config["corrected_source_id"] if change_type == "added" else config["original_source_id"]
                )
                if record["revision_source_id"] != expected_source_id or record["source_id"] != expected_source_id:
                    raise SpendingDataError("revision-diff row source does not match its change type")
                semantic_record = _typed_csv_record(record)
                expected_fingerprint = hashlib.sha256(_canonical_json(_semantic_row(semantic_record))).hexdigest()
                if record["row_fingerprint"] != expected_fingerprint:
                    raise SpendingDataError("revision-diff row fingerprint does not match its analytical values")
                diff_counts[record["change_type"]] += 1
                typed_diff_rows.append(semantic_record)
        if diff_counts["removed"] != comparison["removed_rows"]:
            raise SpendingDataError("revision-diff removed-row count does not match its summary")
        if diff_counts["added"] != comparison["added_rows"]:
            raise SpendingDataError("revision-diff added-row count does not match its summary")
        if verify_against_sources:
            expected_diff_rows, expected_revision = _revision_diff(source, parquet_records, parquet_records)
            if revision != expected_revision:
                raise SpendingDataError("revision summary does not match the registered source workbooks")
            if typed_diff_rows != expected_diff_rows:
                raise SpendingDataError("revision-diff rows do not match the registered source workbooks")


def validate_top_level_finances() -> None:
    registry = _load_yaml(FINANCE_REGISTRY)
    if registry.get("schema_version") != 1:
        raise SpendingDataError("unsupported top-level-finance schema")
    years = registry.get("financial_years")
    if not isinstance(years, list):
        raise SpendingDataError("top-level finance registry must contain financial_years")
    source_ids = {source["id"] for source in registry.get("sources", [])}
    for year in years:
        figures = year.get("figures", [])
        by_id = {figure["id"]: figure for figure in figures}
        if len(by_id) != len(figures):
            raise SpendingDataError(f"duplicate figure ID in {year.get('financial_year')}")
        for figure in figures:
            if int(figure["amount_gbp"]) < 0:
                raise SpendingDataError(f"{figure['id']} has a negative headline amount")
            if "source_id" in figure and figure["source_id"] not in source_ids:
                raise SpendingDataError(f"{figure['id']} refers to unknown source {figure['source_id']}")
            if "derived_from" not in figure:
                continue
            parts = [by_id[figure_id] for figure_id in figure["derived_from"]]
            if len(parts) != 2:
                raise SpendingDataError(f"{figure['id']} must derive from exactly two figures")
            if parts[0]["boundary"] != parts[1]["boundary"] or parts[0]["boundary"] != figure["boundary"]:
                raise SpendingDataError(f"{figure['id']} mixes accounting boundaries")
            expected = int(parts[0]["amount_gbp"]) - int(parts[1]["amount_gbp"])
            if int(figure["amount_gbp"]) != expected:
                raise SpendingDataError(f"{figure['id']} is {figure['amount_gbp']}; expected {expected}")


# Coverage is asserted explicitly per source, not inferred from period strings or
# workbook counts, so an incomplete period can never silently read as a full year.
# "months" is the count of calendar months the release's rows can fall in; it is
# None for the fiscal-year release, which does not map onto calendar-year bars.
SPEND_SUMMARY_COVERAGE: dict[str, dict[str, Any]] = {
    "2013-14-payments-over-25000": {"kind": "fiscal_year", "months": None, "label": "2013/14"},
    "2014-payments-over-25000": {"kind": "partial_calendar_year", "months": 9, "label": "2014"},
    "2015-payments-over-25000": {"kind": "full_calendar_year", "months": 12, "label": "2015"},
    "2016-payments-over-25000": {"kind": "full_calendar_year", "months": 12, "label": "2016"},
    "2017-payments-over-25000": {"kind": "full_calendar_year", "months": 12, "label": "2017"},
    "2018-payments-over-25000": {"kind": "full_calendar_year", "months": 12, "label": "2018"},
    "2019-payments-over-25000": {"kind": "full_calendar_year", "months": 12, "label": "2019"},
    "2020-payments-over-25000": {"kind": "full_calendar_year", "months": 12, "label": "2020"},
    "2021-payments-over-25000": {"kind": "full_calendar_year", "months": 12, "label": "2021"},
    "2022-payments-over-25000": {"kind": "full_calendar_year", "months": 12, "label": "2022"},
    "2023-payments-over-25000": {"kind": "full_calendar_year", "months": 12, "label": "2023"},
    "2024-payments-over-25000": {"kind": "full_calendar_year", "months": 12, "label": "2024"},
    "2025-09-payments-over-25000": {"kind": "partial_calendar_year", "months": 1, "label": "2025 (Sep only)"},
}


def build_spend_summary(target: Path | None = None) -> Path:
    registry = _load_yaml(SOURCE_REGISTRY)
    sources = registry.get("sources", [])
    registered_ids = {str(source["id"]) for source in sources}
    if set(SPEND_SUMMARY_COVERAGE) - registered_ids:
        raise SpendingDataError("spend summary coverage refers to an unregistered source")

    years: list[dict[str, Any]] = []
    for source in sources:
        source_id = str(source["id"])
        coverage = SPEND_SUMMARY_COVERAGE.get(source_id)
        if coverage is None:
            continue
        source_obj = load_source(source_id)
        release = release_path(source_obj)
        validation_path = release / "validation.json"
        if not validation_path.exists():
            raise SpendingDataError(f"no committed release found for {source_id}; build it first")
        validation = json.loads(validation_path.read_text(encoding="utf-8"))
        stats = validation["statistics"]
        years.append(
            {
                "source_id": source_id,
                "period": str(source_obj["period"]),
                "label": coverage["label"],
                "coverage": coverage["kind"],
                "months_covered": coverage["months"],
                "rows": stats["rows"],
                "total_gbp": str(Decimal(stats["total_gbp"])),
                "nhs_england_total_gbp": str(Decimal(stats["nhs_england_total_gbp"])),
            }
        )
    years.sort(key=lambda year: year["period"])

    summary = {
        "schema_version": 1,
        "generated_from": sorted(SPEND_SUMMARY_COVERAGE),
        "years": years,
    }
    target = target or (ROOT / "spending" / "spend-by-year.json")
    target.write_bytes(_json_bytes(summary))
    return target


def build_supplier_spend_summary(target: Path | None = None) -> Path:
    registry = _load_yaml(SOURCE_REGISTRY)
    sources = registry.get("sources", [])
    registered_ids = {str(source["id"]) for source in sources}
    if set(SPEND_SUMMARY_COVERAGE) - registered_ids:
        raise SpendingDataError("spend summary coverage refers to an unregistered source")

    periods: list[dict[str, Any]] = []
    suppliers: dict[str, dict[str, Any]] = {}
    for source_id in sorted(SPEND_SUMMARY_COVERAGE):
        coverage = SPEND_SUMMARY_COVERAGE[source_id]
        source_obj = load_source(source_id)
        period = str(source_obj["period"])
        parquet_path = release_path(source_obj) / "payments.parquet"
        if not parquet_path.exists():
            raise SpendingDataError(f"no committed release found for {source_id}; build it first")
        periods.append(
            {
                "period": period,
                "source_id": source_id,
                "label": coverage["label"],
                "coverage": coverage["kind"],
                "months_covered": coverage["months"],
            }
        )
        connection = duckdb.connect(":memory:")
        try:
            rows = connection.execute(
                """
                SELECT supplier, count(*) AS row_count, sum(amount_gbp) AS total_gbp,
                       sum(CASE WHEN entity = 'NHS England' THEN amount_gbp ELSE 0 END) AS nhs_england_total_gbp
                FROM read_parquet(?)
                GROUP BY supplier
                """,
                [str(parquet_path)],
            ).fetchall()
        finally:
            connection.close()
        for supplier, row_count, total_gbp, nhs_england_total_gbp in rows:
            entry = suppliers.setdefault(supplier, {"total_rows": 0, "total_gbp": Decimal("0.00"), "years": []})
            entry["total_rows"] += row_count
            entry["total_gbp"] += total_gbp
            entry["years"].append(
                {
                    "period": period,
                    "rows": row_count,
                    "total_gbp": str(total_gbp),
                    "nhs_england_total_gbp": str(nhs_england_total_gbp),
                }
            )

    supplier_list = []
    for supplier, entry in sorted(suppliers.items()):
        entry["years"].sort(key=lambda year: year["period"])
        supplier_list.append(
            {
                "supplier": supplier,
                "total_rows": entry["total_rows"],
                "total_gbp": str(entry["total_gbp"]),
                "years": entry["years"],
            }
        )
    periods.sort(key=lambda item: item["period"])

    summary = {
        "schema_version": 1,
        "generated_from": sorted(SPEND_SUMMARY_COVERAGE),
        "periods": periods,
        "suppliers": supplier_list,
    }
    target = target or (ROOT / "spending" / "supplier-spend.json")
    target.write_bytes(_json_bytes(summary))
    return target


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    procurement_fetch_parser = subparsers.add_parser(
        "procurement-fetch", help="download and verify one registered procurement source"
    )
    procurement_fetch_parser.add_argument("--source", required=True, help="procurement source registry ID")
    subparsers.add_parser("procurement-build", help="build and validate the procurement fixture release")
    subparsers.add_parser("procurement-validate", help="validate the procurement fixture release")
    inventory_parser = subparsers.add_parser(
        "inventory", help="discover and classify the historical payments corpus"
    )
    inventory_parser.add_argument(
        "--observed-on", required=True, type=dt.date.fromisoformat, help="inventory observation date (YYYY-MM-DD)"
    )
    inventory_parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "spending" / "historical-schema-census.json",
        help="machine-readable census output",
    )
    inventory_parser.add_argument(
        "--download-directory",
        type=Path,
        default=SOURCE_DIR / "inventory",
        help="ignored raw-source cache",
    )
    fetch_parser = subparsers.add_parser("fetch", help="download and verify the registered source workbook")
    fetch_parser.add_argument("--source", default="2025-09-payments-over-25000", help="source registry ID")
    fetch_parser.add_argument("--archive", action="store_true", help="download the archived source URL")
    for command, help_text in (
        ("build", "build and validate the local data release"),
        ("validate", "validate source, finance registry, and built release"),
        ("all", "fetch, build, and validate"),
    ):
        command_parser = subparsers.add_parser(command, help=help_text)
        command_parser.add_argument("--source", default="2025-09-payments-over-25000", help="source registry ID")
        if command == "all":
            command_parser.add_argument("--archive", action="store_true", help="download the archived source URL")
    subparsers.add_parser(
        "spend-summary", help="rebuild the committed spend-by-year and per-supplier chart data"
    )
    nhsd_fetch_parser = subparsers.add_parser(
        "nhsd-fetch", help="download and verify one registered NHS Digital/HSCIC spend-information source"
    )
    nhsd_fetch_parser.add_argument("--source", required=True, help="NHSD source registry ID")
    for command, help_text in (
        ("nhsd-build", "build and validate the local NHSD release"),
        ("nhsd-validate", "validate the built NHSD release against its registry entry"),
        ("nhsd-all", "fetch, build, and validate one registered NHSD source"),
    ):
        nhsd_command_parser = subparsers.add_parser(command, help=help_text)
        nhsd_command_parser.add_argument("--source", required=True, help="NHSD source registry ID")
    subparsers.add_parser("nhsd-summary", help="rebuild the committed NHSD spend-by-year series")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        if args.command.startswith("procurement-"):
            registry = nhse_procurement.load_registry()
            if args.command == "procurement-fetch":
                path = nhse_procurement.fetch_source(args.source, registry)
                print(f"verified procurement source: {path}", file=sys.stderr)
            elif args.command == "procurement-build":
                target = nhse_procurement.build_release()
                nhse_procurement.validate_release(target, registry, verify_against_sources=True)
                print(f"built procurement release: {target}", file=sys.stderr)
            else:
                target = nhse_procurement.release_path(registry)
                nhse_procurement.validate_release(
                    target,
                    registry,
                    verify_against_sources=True,
                )
                print(f"validated procurement release: {target}", file=sys.stderr)
            return 0
        if args.command == "inventory":
            result = build_inventory(args.output, args.download_directory, args.observed_on)
            print(
                f"wrote inventory: {args.output} ({result['file_count']} files, "
                f"{result['schema_count']} physical schemas, {result['column_profile_count']} column profiles)",
                file=sys.stderr,
            )
            return 0
        if args.command == "spend-summary":
            target = build_spend_summary()
            print(f"wrote spend summary: {target}", file=sys.stderr)
            supplier_target = build_supplier_spend_summary()
            print(f"wrote supplier spend summary: {supplier_target}", file=sys.stderr)
            return 0
        if args.command == "nhsd-summary":
            target = nhse_nhsd_spending.build_nhsd_spend_summary()
            print(f"wrote NHSD spend summary: {target}", file=sys.stderr)
            return 0
        if args.command.startswith("nhsd-"):
            nhsd_registry = nhse_nhsd_spending.load_registry()
            nhsd_source = nhse_nhsd_spending.load_source(args.source, nhsd_registry)
            if args.command in {"nhsd-fetch", "nhsd-all"}:
                path = nhse_nhsd_spending.fetch_source(nhsd_source)
                print(f"verified NHSD source: {path}", file=sys.stderr)
            else:
                path = nhse_nhsd_spending.verify_registered_source(nhsd_source)
            if args.command in {"nhsd-build", "nhsd-all"}:
                target = nhse_nhsd_spending.build_release(path, nhsd_source, nhse_nhsd_spending.release_path(nhsd_source))
                print(f"built NHSD release: {target}", file=sys.stderr)
            if args.command in {"nhsd-validate", "nhsd-all"}:
                nhse_nhsd_spending.validate_release(
                    nhse_nhsd_spending.release_path(nhsd_source), nhsd_source, verify_against_source=True, source_file=path
                )
                print(f"validated NHSD release: {nhse_nhsd_spending.release_path(nhsd_source)}", file=sys.stderr)
            return 0
        source = load_source(args.source)
        if args.command in {"fetch", "all"}:
            paths = fetch_sources(source, use_archive=getattr(args, "archive", False))
            for path in paths:
                print(f"verified source: {path}", file=sys.stderr)
        else:
            paths = verify_registered_sources(source)

        if args.command in {"build", "all"}:
            target = build_release(paths, source, release_path(source))
            print(f"built release: {target}", file=sys.stderr)
        if args.command in {"validate", "all"}:
            validate_top_level_finances()
            validate_release(release_path(source), source, verify_against_sources=True)
            print(f"validated release: {release_path(source)}", file=sys.stderr)
        return 0
    except (
        OSError,
        SpendingDataError,
        duckdb.Error,
        yaml.YAMLError,
        InventoryError,
        nhse_procurement.ProcurementDataError,
        nhse_nhsd_spending.NhsdDataError,
    ) as error:
        print(f"error: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
