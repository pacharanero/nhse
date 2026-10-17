#!/usr/bin/env python3
"""Discover and classify the NHS England payments-over-£25,000 corpus."""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import tempfile
import urllib.request
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlparse

import openpyxl


LANDING_PAGE_URL = "https://www.england.nhs.uk/contact-us/pub-scheme/spend/"
USER_AGENT = "nhse-research-data/0.1 (+https://github.com/pacharanero/nhse)"
MONTHS = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "ocotober": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}


class InventoryError(RuntimeError):
    """A source inventory could not be discovered or inspected safely."""


class _LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag != "a":
            return
        href = dict(attrs).get("href")
        if href:
            self.links.append(href)


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _canonical_json(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode("utf-8")


def _fetch(url: str) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=120) as response:
        return response.read()


def _links(document: bytes, base_url: str) -> list[str]:
    parser = _LinkParser()
    parser.feed(document.decode("utf-8"))
    return list(dict.fromkeys(urljoin(base_url, link) for link in parser.links))


def _collection_year(url: str) -> int:
    match = re.search(r"payments-over-25k-reports?-(20\d{2})", url, re.IGNORECASE)
    if not match:
        raise InventoryError(f"cannot determine collection year from {url}")
    return int(match.group(1))


def _month_from_filename(filename: str) -> int | None:
    words = re.findall(r"[a-z]+", filename.casefold())
    matches = [MONTHS[word] for word in words if word in MONTHS]
    return matches[0] if matches else None


def _cache_path(directory: Path, url: str) -> Path:
    filename = Path(urlparse(url).path).name
    return directory / f"{_sha256_bytes(url.encode('utf-8'))[:16]}-{filename}"


def _download(url: str, target: Path, *, refresh: bool = False) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists() and not refresh:
        return target
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(prefix=f".{target.name}.", suffix=".part", dir=target.parent, delete=False) as output:
            temporary = Path(output.name)
            request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            with urllib.request.urlopen(request, timeout=120) as response:
                shutil.copyfileobj(response, output, length=1024 * 1024)
            output.flush()
            os.fsync(output.fileno())
        os.replace(temporary, target)
        return target
    except Exception:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        raise


def _header_candidate(rows: list[tuple[int, list[Any]]]) -> tuple[int | None, list[str | None]]:
    for row_number, values in rows:
        normalized = [str(value).strip().casefold() if value is not None else "" for value in values]
        if "entity" in normalized and "supplier" in normalized and any("amount" in value for value in normalized):
            return row_number, [None if value is None else str(value).strip() for value in values]
    if rows:
        row_number, values = rows[0]
        return row_number, [None if value is None else str(value).strip() for value in values]
    return None, []


def _inspect_xlsx(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheets: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            populated: list[tuple[int, list[Any]]] = []
            for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30)), start=1):
                values = [cell.value for cell in row]
                while values and values[-1] is None:
                    values.pop()
                if any(value is not None and (not isinstance(value, str) or value.strip()) for value in values):
                    populated.append((row_number, values))
            header_row, headers = _header_candidate(populated)
            sheets.append(
                {
                    "name": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "header_row": header_row,
                    "headers": headers,
                }
            )
    finally:
        workbook.close()
    return {"format": "xlsx", "sheets": sheets}


def _inspect_csv(path: Path) -> dict[str, Any]:
    raw = path.read_bytes()
    encoding = "utf-8-sig"
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError:
        encoding = "windows-1252"
        text = raw.decode(encoding)
    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    rows = []
    for row_number, row in enumerate(csv.reader(text.splitlines(), delimiter=delimiter), start=1):
        if any(value.strip() for value in row):
            rows.append((row_number, row))
        if row_number >= 30:
            break
    header_row, headers = _header_candidate(rows)
    return {
        "format": "csv",
        "encoding": encoding,
        "delimiter": delimiter,
        "header_row": header_row,
        "headers": headers,
    }


def _inspect_file(path: Path) -> dict[str, Any]:
    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        return _inspect_xlsx(path)
    if suffix == ".csv":
        return _inspect_csv(path)
    return {"format": suffix.removeprefix("."), "status": "unsupported"}


def _schema_projection(inspection: dict[str, Any]) -> dict[str, Any]:
    if inspection["format"] == "xlsx":
        return {
            "format": "xlsx",
            "sheets": [
                {"name": sheet["name"], "header_row": sheet["header_row"], "headers": sheet["headers"]}
                for sheet in inspection["sheets"]
            ],
        }
    return {
        key: inspection[key]
        for key in ("format", "encoding", "delimiter", "header_row", "headers")
        if key in inspection
    }


def _column_projection(inspection: dict[str, Any]) -> dict[str, Any]:
    if inspection["format"] == "xlsx":
        return {
            "format": "xlsx",
            "header_sets": [sheet["headers"] for sheet in inspection["sheets"] if sheet["headers"]],
        }
    return {
        key: inspection[key]
        for key in ("format", "encoding", "delimiter", "headers")
        if key in inspection
    }


def build_inventory(output: Path, download_directory: Path, observed_on: dt.date) -> dict[str, Any]:
    landing_document = _fetch(LANDING_PAGE_URL)
    landing_links = _links(landing_document, LANDING_PAGE_URL)
    collection_urls = sorted(
        {
            url
            for url in landing_links
            if re.search(r"/publication/payments-over-25k-reports?-20\d{2}/?$", url, re.IGNORECASE)
        },
        key=_collection_year,
    )
    legacy_urls = sorted(
        {
            url
            for url in landing_links
            if "1314" in url.casefold() and urlparse(url).path.casefold().endswith((".xlsx", ".csv"))
        }
    )
    if not collection_urls or not legacy_urls:
        raise InventoryError("canonical landing page did not expose the expected collection and 2013/14 links")

    collections: list[dict[str, Any]] = []
    files: list[dict[str, Any]] = []
    downloaded_paths: dict[str, Path] = {}
    for year, collection_url, document_urls in [
        (2013, LANDING_PAGE_URL, legacy_urls),
        *[
            (
                _collection_year(collection_url),
                collection_url,
                None,
            )
            for collection_url in collection_urls
        ],
    ]:
        if document_urls is None:
            collection_document = _fetch(collection_url)
            document_urls = [
                url
                for url in _links(collection_document, collection_url)
                if urlparse(url).path.casefold().endswith((".xlsx", ".xls", ".csv"))
            ]
            collection_digest = _sha256_bytes(collection_document)
            collection_id = str(year)
        else:
            collection_digest = _sha256_bytes(landing_document)
            collection_id = "2013-14"
        document_urls = list(dict.fromkeys(document_urls))
        observed_months: list[int] = []
        collection_file_indexes: list[int] = []
        for source_url in document_urls:
            filename = Path(urlparse(source_url).path).name
            month = _month_from_filename(filename)
            if month is not None and collection_id != "2013-14":
                observed_months.append(month)
            if source_url not in downloaded_paths:
                downloaded_paths[source_url] = _download(
                    source_url,
                    _cache_path(download_directory, source_url),
                    refresh=True,
                )
            local_path = downloaded_paths[source_url]
            inspection = _inspect_file(local_path)
            schema_projection = _schema_projection(inspection)
            schema_sha256 = _sha256_bytes(_canonical_json(schema_projection))
            column_projection = _column_projection(inspection)
            column_profile_sha256 = _sha256_bytes(_canonical_json(column_projection))
            entry = {
                "collection_id": collection_id,
                "period": "2013-14"
                if "1314" in filename.casefold()
                else f"{year:04d}-{month:02d}"
                if month is not None and collection_id != "2013-14"
                else collection_id,
                "source_url": source_url,
                "filename": filename,
                "bytes": local_path.stat().st_size,
                "sha256": hashlib.sha256(local_path.read_bytes()).hexdigest(),
                "revision_marker": "corrected"
                if "corrected" in filename.casefold()
                else "revised"
                if "revised" in filename.casefold()
                else "v2"
                if "v2" in filename.casefold()
                else None,
                "schema_sha256": schema_sha256,
                "column_profile_sha256": column_profile_sha256,
                "inspection": inspection,
            }
            collection_file_indexes.append(len(files))
            files.append(entry)
        observed = sorted(set(observed_months))
        collections.append(
            {
                "id": collection_id,
                "collection_url": collection_url,
                "collection_page_sha256": collection_digest,
                "file_count": len(document_urls),
                "observed_months": observed,
                "missing_months": [month for month in range(1, 13) if month not in observed]
                if collection_id != "2013-14"
                else [],
                "file_indexes": collection_file_indexes,
            }
        )

    schema_counts = Counter(file["schema_sha256"] for file in files)
    schema_ids = {digest: f"schema-{index:02d}" for index, digest in enumerate(sorted(schema_counts), start=1)}
    schemas = []
    for digest in sorted(schema_counts):
        matching = [file for file in files if file["schema_sha256"] == digest]
        schemas.append(
            {
                "id": schema_ids[digest],
                "sha256": digest,
                "file_count": len(matching),
                "collections": sorted({file["collection_id"] for file in matching}),
                "projection": _schema_projection(matching[0]["inspection"]),
            }
        )
    for file in files:
        file["schema_id"] = schema_ids[file.pop("schema_sha256")]

    column_profile_counts = Counter(file["column_profile_sha256"] for file in files)
    column_profile_ids = {
        digest: f"column-profile-{index:02d}"
        for index, digest in enumerate(sorted(column_profile_counts), start=1)
    }
    column_profiles = []
    for digest in sorted(column_profile_counts):
        matching = [file for file in files if file["column_profile_sha256"] == digest]
        column_profiles.append(
            {
                "id": column_profile_ids[digest],
                "sha256": digest,
                "file_count": len(matching),
                "collections": sorted({file["collection_id"] for file in matching}),
                "projection": _column_projection(matching[0]["inspection"]),
            }
        )
    for file in files:
        file["column_profile_id"] = column_profile_ids[file.pop("column_profile_sha256")]

    result = {
        "schema_version": 1,
        "observed_on": observed_on.isoformat(),
        "landing_page_url": LANDING_PAGE_URL,
        "landing_page_sha256": _sha256_bytes(landing_document),
        "collection_count": len(collections),
        "file_count": len(files),
        "unique_source_count": len({file["source_url"] for file in files}),
        "file_count_definition": "Logical link occurrences; one source can appear in more than one collection.",
        "schema_count": len(schemas),
        "column_profile_count": len(column_profiles),
        "collections": collections,
        "schemas": schemas,
        "column_profiles": column_profiles,
        "files": files,
        "limitations": [
            "This census records links exposed by the canonical landing and collection pages on the observation date; absence is not proof that a source never existed.",
            "Schema fingerprints cover workbook sheet names and detected header rows, or CSV encoding, delimiter and detected headers; they do not assert semantic equivalence.",
            "Revision markers are filename observations only; superseded bytes require separate archival evidence and comparison.",
            "The 2013/14 Q4 CSV is linked from both the canonical landing page and the 2014 collection; both logical occurrences are preserved.",
        ],
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.name}.tmp")
    temporary.write_text(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    os.replace(temporary, output)
    return result
